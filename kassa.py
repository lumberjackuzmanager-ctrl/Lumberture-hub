# -*- coding: utf-8 -*-
"""Касса (вместо ежемесячных Excel «касса <Месяц> <Год>.xlsx»): один раздел на все месяцы.

Как в Excel:
  • день: кирим (1С и прочий), чиқим (лист/кромка учун тўлов, ойлик-аванс, харажатлар, шериклар), кун боши/охири, факт, фарқ;
  • месяц («от»): итоги по поставщикам, ойлик (ойлик / бошига / аванс / охирига), харажатлар, кирим-чиқим, қолдиқ;
  • табель и ойлик («т»), KPI (распил/кромка, кундуз/тун × ставка), овқат («х»).
Остатки переходят из месяца в месяц сами: кун боши = факт прошлого дня, ой бошига = ой охирига прошлого месяца.
Аренда в кассу берётся из платформы (оплаты арендаторов), в Excel её не вводят.
"""
import time, calendar, os, re, json
from collections import defaultdict
from datetime import date

from flask import request, render_template, redirect, url_for, flash, abort, jsonify

A = None
MONTHS = ["", "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]
KINDS = {"sheet": "Лист учун тўлов", "edge": "Кромка учун тўлов", "exp": "Харажатлар", "owner": "Шериклар"}
KPI_KINDS = {"raspil": "Распил", "kromka": "Кромка"}
KPI_RESYNC = []          # автомат ҳисоблар (цех, Базис) — «автоматга қайтариш»да қайта ҳисоблайди: fn(ym)
# харажатлар бўлимлари (кунлик кассада ва ой итогида алоҳида)
EXP_GROUPS = [("food", "Овқат ва хўжалик"), ("machine", "Станоклар"), ("basis", "Базис"), ("tax", "Бухгалтер / налог"),
              ("office", "Компьютер ва канцтовар"), ("delivery", "Доставка ва шофёрлар"), ("rent", "Аренда ва коммунал"), ("other", "Бошқа")]
EXP_GROUP_NAMES = dict(EXP_GROUPS)

def exp_group_guess(name):
    n = (name or "").lower()
    if n.startswith("аренда") or "комунал" in n or "коммунал" in n:
        return "rent"
    if any(w in n for w in ("арра", "кромка nb", "альтендорф", "npc", "nd512", "присадка", "ровер", "брантд", "компрессор", "овал", "станок")):
        return "machine"
    if "доставка" in n or re.search(r"ака\s*\d{3}\b", n):         # «Отабек ака 164», «Равшан ака 842» — шофёрлар
        return "delivery"
    if "базис" in n:
        return "basis"
    if "бухгалтер" in n or "налог" in n:
        return "tax"
    if any(w in n for w in ("компьютер", "компютер", "канц", "констовар", "конставар", "интернет", "телефон", "копга", "вытишка")):
        return "office"
    if any(w in n for w in ("завтрик", "обед", "ужин", "продукт", "куш-харажат", "хўжалик")):
        return "food"
    return "other"
SHIFTS = {"d": "кундуз", "n": "тун"}
TS_OFF = ("х", "x", "")

SCHEMA = """
CREATE TABLE IF NOT EXISTS k_items(id INTEGER PRIMARY KEY, kind TEXT NOT NULL, name TEXT NOT NULL, sort INTEGER NOT NULL DEFAULT 100,
  active INTEGER NOT NULL DEFAULT 1, share REAL NOT NULL DEFAULT 0, UNIQUE(kind, name));
CREATE TABLE IF NOT EXISTS k_workers(id INTEGER PRIMARY KEY, name TEXT NOT NULL UNIQUE, salary REAL NOT NULL DEFAULT 0,
  cur TEXT NOT NULL DEFAULT 'UZS', sort INTEGER NOT NULL DEFAULT 100, active INTEGER NOT NULL DEFAULT 1);
CREATE TABLE IF NOT EXISTS k_months(ym TEXT PRIMARY KEY, open_uzs REAL NOT NULL DEFAULT 0, open_usd REAL NOT NULL DEFAULT 0,
  open_set INTEGER NOT NULL DEFAULT 0, onec_usd REAL, meal_b REAL NOT NULL DEFAULT 12000, meal_l REAL NOT NULL DEFAULT 17000,
  meal_d REAL NOT NULL DEFAULT 17000, meal_paid REAL NOT NULL DEFAULT 0, note TEXT DEFAULT '', source TEXT DEFAULT '');
CREATE TABLE IF NOT EXISTS k_days(ym TEXT NOT NULL, day INTEGER NOT NULL, rate REAL, in1_uzs REAL NOT NULL DEFAULT 0, in1_usd REAL NOT NULL DEFAULT 0,
  in2_uzs REAL NOT NULL DEFAULT 0, in2_usd REAL NOT NULL DEFAULT 0, fact_uzs REAL NOT NULL DEFAULT 0, fact_usd REAL NOT NULL DEFAULT 0,
  has_fact INTEGER NOT NULL DEFAULT 0, note TEXT DEFAULT '', updated_at TEXT, PRIMARY KEY(ym, day));
CREATE TABLE IF NOT EXISTS k_cells(ym TEXT NOT NULL, day INTEGER NOT NULL, kind TEXT NOT NULL, ref_id INTEGER NOT NULL,
  uzs REAL NOT NULL DEFAULT 0, usd REAL NOT NULL DEFAULT 0, expr_uzs TEXT DEFAULT '', expr_usd TEXT DEFAULT '',
  PRIMARY KEY(ym, day, kind, ref_id));
CREATE TABLE IF NOT EXISTS k_ws(ym TEXT NOT NULL, worker_id INTEGER NOT NULL, sort INTEGER NOT NULL DEFAULT 100, salary REAL NOT NULL DEFAULT 0,
  days_override REAL, extra REAL NOT NULL DEFAULT 0, total_override REAL, opening REAL NOT NULL DEFAULT 0, opening_set INTEGER NOT NULL DEFAULT 0,
  ts TEXT NOT NULL DEFAULT '', PRIMARY KEY(ym, worker_id));
CREATE TABLE IF NOT EXISTS k_kpi_cols(id INTEGER PRIMARY KEY, worker_id INTEGER NOT NULL, kind TEXT NOT NULL, shift TEXT NOT NULL,
  rate REAL NOT NULL DEFAULT 0, sort INTEGER NOT NULL DEFAULT 100, active INTEGER NOT NULL DEFAULT 1, UNIQUE(worker_id, kind, shift));
CREATE TABLE IF NOT EXISTS k_kpi_mcols(ym TEXT NOT NULL, col_id INTEGER NOT NULL, rate REAL NOT NULL DEFAULT 0, sort INTEGER NOT NULL DEFAULT 100,
  xl_col TEXT DEFAULT '', PRIMARY KEY(ym, col_id));
CREATE TABLE IF NOT EXISTS k_kpi(ym TEXT NOT NULL, day INTEGER NOT NULL, col_id INTEGER NOT NULL, qty REAL NOT NULL DEFAULT 0,
  PRIMARY KEY(ym, day, col_id));
CREATE TABLE IF NOT EXISTS k_kpi_lock(ym TEXT NOT NULL, day INTEGER NOT NULL, col_id INTEGER NOT NULL, at TEXT DEFAULT '',
  PRIMARY KEY(ym, day, col_id));
CREATE TABLE IF NOT EXISTS k_meals(ym TEXT NOT NULL, day INTEGER NOT NULL, b REAL NOT NULL DEFAULT 0, l REAL NOT NULL DEFAULT 0,
  d REAL NOT NULL DEFAULT 0, PRIMARY KEY(ym, day));
"""

DEFAULT_ITEMS = {
    "sheet": ["Абдулло г-с", "Албери", "База Эман", "Ибн-Сино", "Каштан", "Келес", "Кастомону", "Панелный", "Сергили", "ТАПОиЧ",
              "Timber zone", "Ёш Ленинчи", "Назарбек", "Нодир ака", "Ламарти", "Гиламсарой", "Бозор"],
    "edge": ["Абдумалик", "Алкан", "Бозор", "Клоп", "Хенкел", "Чинобод"],
    "exp": ["Завтрик", "Обед", "Ужин", "Продукты", "куш-харажат", "Бухгалтер/Налог", "Арра чарх", "Кромка NB536", "Альтендорф WA8",
            "NPC 330 П.центр", "ND512DT Присадка", "Ровер", "Брантд 120", "Компрессор", "Базис", "Овал кромка", "Доставка", "Комуналка", "Аренда", "Компьютер ва канцтовар"],
    "owner": [("Жамшид", 5), ("Лочин", 3), ("Жасур", 2)],
}


# ---------------------------------------------------------------- helpers
def ym_of(y, m):
    return f"{y:04d}-{m:02d}"

def ym_parse(ym):
    try:
        y, m = (int(x) for x in str(ym).split("-")[:2])
        if 2000 <= y <= 2100 and 1 <= m <= 12:
            return y, m
    except (TypeError, ValueError):
        pass
    return None

def ym_shift(ym, k):
    y, m = ym_parse(ym)
    i = y * 12 + (m - 1) + k
    return ym_of(i // 12, i % 12 + 1)

def ym_days(ym):
    y, m = ym_parse(ym)
    return calendar.monthrange(y, m)[1]

def ym_title(ym):
    y, m = ym_parse(ym)
    return f"{MONTHS[m]} {y}"

def ym_today():
    t = date.today()
    return ym_of(t.year, t.month)

_EXPR_OK = re.compile(r"^[0-9+\-*/(). ]+$")

def calc(expr):
    """«150000+120000+250000» → 520000 (как формула в Excel). Пусто → 0. Ошибка → ValueError."""
    s = str(expr or "").strip().lstrip("=").replace("\u00a0", "").replace("\u202f", "").replace("\u2212", "-").replace("×", "*")
    s = re.sub(r"(?<=\d) (?=\d)", "", s).replace(",", ".")
    if not s:
        return 0.0
    if not _EXPR_OK.match(s) or "**" in s or "//" in s:
        raise ValueError(expr)
    try:
        v = eval(compile(s, "<calc>", "eval"), {"__builtins__": {}}, {})
    except Exception:
        raise ValueError(expr)
    return float(v)

def kin(v):
    """Число для поля ввода: 13556292.55 (без «e+07» и лишних нулей)."""
    try:
        v = float(v or 0)
    except (TypeError, ValueError):
        return v
    t = f"{v:.2f}".rstrip("0").rstrip(".")
    return "0" if t in ("-0", "") else t

def fnum(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0

def ts_list(ts, n=31):
    out = (ts or "").split("|")
    return (out + [""] * n)[:n]

def ts_days(ts, hours=9.0):
    """Табель: число = часы (9 = полный день), «х» — выходной, буква — отметка без часов."""
    h, marks = 0.0, 0
    for c in ts_list(ts):
        c = str(c).strip().replace(",", ".")
        if c.lower() in TS_OFF:
            continue
        try:
            h += float(c)
        except ValueError:
            marks += 1
    return h / hours if hours else 0.0, marks


# ---------------------------------------------------------------- model
def init_tables():
    con = A.connect()
    con.executescript(SCHEMA)
    if "note_uzs" not in [r[1] for r in con.execute("PRAGMA table_info(k_cells)")]:
        con.execute("ALTER TABLE k_cells ADD COLUMN note_uzs TEXT NOT NULL DEFAULT ''")    # Excel'дагидек катак изоҳи (примечание)
        con.execute("ALTER TABLE k_cells ADD COLUMN note_usd TEXT NOT NULL DEFAULT ''")
    if "fio" not in [r[1] for r in con.execute("PRAGMA table_info(k_workers)")]:
        con.execute("ALTER TABLE k_workers ADD COLUMN fio TEXT DEFAULT ''")      # Фамилия Исм Шарифи (тўлиқ)
    if "fired" not in [r[1] for r in con.execute("PRAGMA table_info(k_workers)")]:
        con.execute("ALTER TABLE k_workers ADD COLUMN fired TEXT")      # ишдан бўшаган сана (YYYY-MM-DD)
    kd = [r[1] for r in con.execute("PRAGMA table_info(k_days)")]
    if "in1_man" not in kd:          # 1 — кассир «Кирим 1С»ни қўлда ўзгартирган (1С'дан автомат янгиланмайди); NULL/0 — 1С'дан олинади
        con.execute("ALTER TABLE k_days ADD COLUMN in1_man INTEGER")
        con.execute("ALTER TABLE k_days ADD COLUMN in2_man INTEGER")
    if "xl_last_day" not in [r[1] for r in con.execute("PRAGMA table_info(k_months)")]:
        con.execute("ALTER TABLE k_months ADD COLUMN xl_last_day INTEGER NOT NULL DEFAULT 0")   # шу кунгача касса Excel'дан (аренда ҳам)
    if "grp" not in [r[1] for r in con.execute("PRAGMA table_info(k_items)")]:
        con.execute("ALTER TABLE k_items ADD COLUMN grp TEXT NOT NULL DEFAULT ''")
        # аренда — битта қатор: плюс тўланган (расход), минус — тушган/қайтган
        paid = con.execute("SELECT id FROM k_items WHERE kind='exp' AND name='Аренда (тўланган)'").fetchone()
        rin = con.execute("SELECT id FROM k_items WHERE kind='exp' AND name='Аренда (тушган, Excel)'").fetchone()
        if paid or rin:
            main = (paid or rin)[0]
            if not con.execute("SELECT 1 FROM k_items WHERE kind='exp' AND name='Аренда' AND id<>?", (main,)).fetchone():
                con.execute("UPDATE k_items SET name='Аренда', active=1 WHERE id=?", (main,))
            if paid and rin:
                for c in con.execute("SELECT ym, day, uzs, usd FROM k_cells WHERE kind='exp' AND ref_id=?", (rin[0],)).fetchall():
                    con.execute("INSERT INTO k_cells(ym,day,kind,ref_id,uzs,usd) VALUES(?,?,'exp',?,?,?) "
                                "ON CONFLICT(ym,day,kind,ref_id) DO UPDATE SET uzs=uzs+excluded.uzs, usd=usd+excluded.usd, expr_uzs='', expr_usd=''",
                                (c[0], c[1], main, c[2], c[3]))
                con.execute("DELETE FROM k_cells WHERE kind='exp' AND ref_id=?", (rin[0],))
                con.execute("DELETE FROM k_items WHERE id=?", (rin[0],))
        for iid, nm in con.execute("SELECT id, name FROM k_items WHERE kind='exp'").fetchall():
            con.execute("UPDATE k_items SET grp=? WHERE id=?", (exp_group_guess(nm), iid))
        if con.execute("SELECT 1 FROM k_items LIMIT 1").fetchone() and \
                not con.execute("SELECT 1 FROM k_items WHERE kind='exp' AND name='Компьютер ва канцтовар'").fetchone():
            mx = con.execute("SELECT COALESCE(MAX(sort),0) FROM k_items WHERE kind='exp'").fetchone()[0]
            con.execute("INSERT INTO k_items(kind,name,sort,active,grp) VALUES('exp','Компьютер ва канцтовар',?,1,'office')", (mx + 10,))
    if "onec_zav" not in [r[1] for r in con.execute("PRAGMA table_info(k_months)")]:
        con.execute("ALTER TABLE k_months ADD COLUMN onec_zav REAL")       # «1С зав.» — 1С бўйича чиқим, солиштириш учун
    if not con.execute("SELECT 1 FROM k_kpi_mcols LIMIT 1").fetchone():
        # KPI устунлари ва ставкалари энди ҳар ой ўзиники (Excel'даги каби)
        con.execute("INSERT OR IGNORE INTO k_kpi_mcols(ym,col_id,rate,sort) SELECT DISTINCT k.ym, c.id, c.rate, c.sort FROM k_kpi k JOIN k_kpi_cols c ON c.id=k.col_id")
    if "svc" not in [r[1] for r in con.execute("PRAGMA table_info(k_workers)")]:
        # ташқи хизмат (бошқа цехга қилдирилган иш): «от» листидаги Ровер / распил ноч / кромка ноч
        con.execute("ALTER TABLE k_workers ADD COLUMN svc INTEGER NOT NULL DEFAULT 0")
        con.execute("UPDATE k_workers SET svc=1 WHERE lower(name) IN ('ровер','распил ноч','кромка ноч')")
    con.executescript("""
        CREATE TABLE IF NOT EXISTS k_noncash(id INTEGER PRIMARY KEY, ym TEXT NOT NULL, day INTEGER NOT NULL, acc_id INTEGER NOT NULL,
          kind TEXT NOT NULL, ref_id INTEGER, uzs REAL NOT NULL DEFAULT 0, usd REAL NOT NULL DEFAULT 0, note TEXT DEFAULT '');
        CREATE INDEX IF NOT EXISTS k_noncash_d ON k_noncash(ym, day);
        CREATE TABLE IF NOT EXISTS k_acc_open(id INTEGER PRIMARY KEY, acc_id INTEGER NOT NULL, date TEXT NOT NULL, uzs REAL DEFAULT 0,
          usd REAL DEFAULT 0, set_by TEXT, set_at TEXT);
    """)
    if not con.execute("SELECT 1 FROM k_items LIMIT 1").fetchone():
        for kind, names in DEFAULT_ITEMS.items():
            for i, n in enumerate(names):
                name, share = (n if isinstance(n, tuple) else (n, 0))
                con.execute("INSERT OR IGNORE INTO k_items(kind,name,sort,share,grp) VALUES(?,?,?,?,?)",
                            (kind, name, (i + 1) * 10, share, exp_group_guess(name) if kind == "exp" else ""))
    con.commit(); con.close()

def items(kind=None, ym=None, all_=False):
    """Строки раздела: активные + те, по которым есть суммы в этом месяце."""
    q = "SELECT * FROM k_items" + (" WHERE kind=?" if kind else "") + " ORDER BY kind, sort, name"
    rows = A.q(q, (kind,) if kind else ())
    if all_:
        return rows
    used = set()
    if ym:
        used = {(r["kind"], r["ref_id"]) for r in A.q("SELECT DISTINCT kind, ref_id FROM k_cells WHERE ym=? AND (uzs<>0 OR usd<>0)", (ym,))}
    return [r for r in rows if r["active"] or (r["kind"], r["id"]) in used]

def month_row(ym):
    return A.q("SELECT * FROM k_months WHERE ym=?", (ym,), one=True)

def ensure_month(ym):
    """Новый месяц: строка месяца + список работников из прошлого месяца (как новый файл из шаблона)."""
    db = A.db()
    if not month_row(ym):
        p = month_row(ym_shift(ym, -1))
        db.execute("INSERT OR IGNORE INTO k_months(ym, meal_b, meal_l, meal_d) VALUES(?,?,?,?)",
                   (ym, p["meal_b"] if p else 12000, p["meal_l"] if p else 17000, p["meal_d"] if p else 17000))
    if not A.q("SELECT 1 FROM k_ws WHERE ym=? LIMIT 1", (ym,)):
        prev = A.q("SELECT s.*, w.salary wsal FROM k_ws s JOIN k_workers w ON w.id=s.worker_id WHERE s.ym=? AND w.active=1 AND (w.fired IS NULL OR w.fired='') ORDER BY s.sort",
                   (ym_shift(ym, -1),))
        if prev:
            for r in prev:
                db.execute("INSERT OR IGNORE INTO k_ws(ym,worker_id,sort,salary) VALUES(?,?,?,?)", (ym, r["worker_id"], r["sort"], r["salary"]))
        else:
            for i, w in enumerate(A.q("SELECT * FROM k_workers WHERE active=1 AND (fired IS NULL OR fired='') ORDER BY sort, name")):
                db.execute("INSERT OR IGNORE INTO k_ws(ym,worker_id,sort,salary) VALUES(?,?,?,?)", (ym, w["id"], (i + 1) * 10, w["salary"]))
        for wid, sal in salary_plan(ym).items():          # олдиндан белгиланган янги оклад (шу ойдан)
            db.execute("UPDATE k_ws SET salary=? WHERE ym=? AND worker_id=?", (sal, ym, wid))
            db.execute("UPDATE k_workers SET salary=? WHERE id=?", (sal, wid))
    db.commit()


def salary_plan(ym):
    """Созлама salary_plan (JSON): [{"worker": id, "from": "ГГГГ-ММ", "salary": сўм}] — янги ой очилганда шу ойдан оклад.
    Эгаси, 25.10: менежер крой — ноябрдан оклад 2 млн + распил бай. Фақат «from» ойи очилганда (кейин қўлда ўзгартирилгани
    кейинги ойларга ўтади). -> {worker: salary}"""
    try:
        plan = json.loads(A.get_setting("salary_plan", "") or "[]")
    except ValueError:
        return {}
    out = {}
    for p in sorted((p for p in plan if isinstance(p, dict)), key=lambda p: str(p.get("from") or "")):
        try:
            if str(p.get("from") or "") == ym:
                out[int(p["worker"])] = float(p["salary"])
        except (KeyError, TypeError, ValueError):
            continue
    return out

def has_data(ym):
    return bool(A.q("SELECT 1 FROM k_days WHERE ym=? LIMIT 1", (ym,)) or A.q("SELECT 1 FROM k_months WHERE ym=? AND open_set=1", (ym,)))

def default_rate(d):
    r = A.q("SELECT rate FROM fx_rates WHERE source<>'cbu' AND COALESCE(rate_date, substr(fetched_at,1,10))<=? "
            "ORDER BY COALESCE(rate_date, substr(fetched_at,1,10)) DESC, id DESC LIMIT 1", (d.isoformat(),), one=True)   # Эман боти; ЦБ эмас
    return r["rate"] if r else (A.fx_rate() or 0)

def rent_by_day(ym):
    """Аренда из платформы: оплаты арендаторов за каждый день месяца (сум, $)."""
    out = defaultdict(lambda: [0.0, 0.0])
    if A.get_setting("kassa_rent", "1") != "1":
        return out
    mr = month_row(ym)
    xl_last = (mr["xl_last_day"] if mr and "xl_last_day" in mr.keys() else 0) or 0   # Excel'дан олинган кунларда аренда Excel'дагидек
    for r in A.q("SELECT pay_date, amount, currency FROM payments WHERE substr(pay_date,1,7)=? AND is_demo=0", (ym,)):
        try:
            d = int(r["pay_date"][8:10])
        except (TypeError, ValueError):
            continue
        if d <= xl_last:
            continue
        out[d][1 if (r["currency"] or "UZS") == "USD" else 0] += r["amount"] or 0
    return out


_memo = {}

def month_calc(ym, _depth=0):
    """Полный расчёт месяца (как лист «от» + дни 1–31). Кэш на время запроса."""
    key = (id(A.db()), ym)
    if key in _memo:
        return _memo[key]
    n = ym_days(ym)
    mr = month_row(ym)
    days = {r["day"]: r for r in A.q("SELECT * FROM k_days WHERE ym=?", (ym,))}
    cells = defaultdict(list)
    for c in A.q("SELECT * FROM k_cells WHERE ym=?", (ym,)):
        cells[c["day"]].append(c)
    moves = defaultdict(lambda: [0.0, 0.0, 0.0, 0.0])      # нақд кассага/дан ўтказма: кирим сўм, $, чиқим сўм, $
    try:
        cash_ids = {r[0] for r in A.q("SELECT id FROM erp_accounts WHERE kind='cash'")}
        for r in A.q("SELECT * FROM erp_moves WHERE deleted=0 AND substr(date,1,7)=?", (ym,)):
            dd, i = int(r["date"][8:10]), (0 if r["cur"] == "uzs" else 1)
            if r["to_acc"] in cash_ids:
                moves[dd][i] += r["amount"]
            if r["from_acc"] in cash_ids:
                moves[dd][i + 2] += r["amount"]
    except Exception:
        pass
    noncash = defaultdict(list)
    try:
        for c in A.q("SELECT * FROM k_noncash WHERE ym=?", (ym,)):
            noncash[c["day"]].append(c)
    except Exception:
        pass
    rent = rent_by_day(ym)
    rates = [r["rate"] for r in days.values() if r["rate"]]
    avg = sum(rates) / len(rates) if rates else (default_rate(date(*ym_parse(ym), 1)) or 12000)

    # остаток на начало месяца
    if mr and mr["open_set"]:
        start = [mr["open_uzs"], mr["open_usd"]]
    elif _depth < 36 and has_data(ym_shift(ym, -1)):
        start = list(month_calc(ym_shift(ym, -1), _depth + 1)["carry"])
    else:
        start = [mr["open_uzs"] if mr else 0, mr["open_usd"] if mr else 0]
    month_start = list(start)

    per_item = defaultdict(lambda: [0.0, 0.0, 0.0])      # (kind,id) -> uzs, usd, $total
    rent_ids = {r["id"] for r in A.q("SELECT id, name FROM k_items WHERE kind='exp'") if r["name"].lower().startswith("аренда")}
    rent_pm = [0.0, 0.0]                                 # аренда: тўланган (+), тушган/қайтган (−), $ да
    per_worker = defaultdict(lambda: [0.0, 0.0, 0.0, 0.0])   # worker -> сум, $, жами $ (кун курси), $ берилгани × кун курси
    dl, tot = [], defaultdict(float)
    carry = list(start)
    for d in range(1, n + 1):
        r = days.get(d)
        rate = (r["rate"] if r and r["rate"] else 0) or avg
        o = dict(day=d, row=r, rate=rate, start=list(carry), rent=list(rent.get(d, [0.0, 0.0])))
        out_uzs = out_usd = 0.0
        g = defaultdict(lambda: [0.0, 0.0])
        for c in cells.get(d, []):
            k = c["kind"]
            if k == "wage":
                pw = per_worker[c["ref_id"]]
                pw[0] += c["uzs"]; pw[1] += c["usd"]; pw[2] += c["uzs"] / rate + c["usd"]; pw[3] += c["usd"] * rate
                out_uzs += c["uzs"]; out_usd += c["usd"]; g["wage"][0] += c["uzs"]; g["wage"][1] += c["usd"]
                continue
            usd_total = c["uzs"] / rate + c["usd"]
            pi = per_item[(k, c["ref_id"])]
            pi[0] += c["uzs"]; pi[1] += c["usd"]; pi[2] += usd_total
            if k == "exp" and c["ref_id"] in rent_ids:
                rent_pm[0 if usd_total > 0 else 1] += usd_total
            if k in ("sheet", "edge"):                     # в Excel поставщики уходят из долларов (сум → $ по курсу дня)
                out_usd += usd_total; g[k][1] += usd_total
            else:
                out_uzs += c["uzs"]; out_usd += c["usd"]; g[k][0] += c["uzs"]; g[k][1] += c["usd"]
        for c in noncash.get(d, []):        # картадан/банкдан тўланган: завод/харажат/ойлик жамига киради, нақд кассага эмас
            k = c["kind"]
            if k == "wage":
                pw = per_worker[c["ref_id"]]
                pw[0] += c["uzs"]; pw[1] += c["usd"]; pw[2] += c["uzs"] / rate + c["usd"]; pw[3] += c["usd"] * rate
            elif k in KINDS:
                pi = per_item[(k, c["ref_id"])]
                pi[0] += c["uzs"]; pi[1] += c["usd"]; pi[2] += c["uzs"] / rate + c["usd"]
        o["noncash"] = noncash.get(d, [])
        mv = moves.get(d, [0.0, 0.0, 0.0, 0.0])
        o["moves"] = mv
        in_uzs = (r["in1_uzs"] + r["in2_uzs"] if r else 0) + o["rent"][0] + mv[0]
        in_usd = (r["in1_usd"] + r["in2_usd"] if r else 0) + o["rent"][1] + mv[1]
        out_uzs += mv[2]; out_usd += mv[3]
        end = [carry[0] + in_uzs - out_uzs, carry[1] + in_usd - out_usd]
        o.update(out=[out_uzs, out_usd], inc=[in_uzs, in_usd], end=end, groups=g, filled=bool(r) or bool(cells.get(d)))
        if r and r["has_fact"]:
            o["fact"] = [r["fact_uzs"], r["fact_usd"]]
            o["diff"] = (end[0] - r["fact_uzs"]) / rate + (end[1] - r["fact_usd"])
            carry = list(o["fact"])
        else:
            o["fact"], o["diff"] = None, None
            carry = list(end)
        if r:
            tot["in1"] += r["in1_uzs"] / rate + r["in1_usd"]
            tot["in2"] += r["in2_uzs"] / rate + r["in2_usd"]
        tot["rent"] += o["rent"][0] / rate + o["rent"][1]
        tot["out"] += out_uzs / rate + out_usd
        if o["diff"] is not None:
            tot["diff"] += o["diff"]
        dl.append(o)

    staff = staff_calc(ym, avg, per_worker, _depth)
    res = dict(ym=ym, n=n, avg=avg, days=dl, start=month_start, carry=carry, per_item=per_item, per_worker=per_worker,
               staff=staff, tot=tot, month=mr)
    # итоги как на листе «от»
    s = lambda kind: sum(v[2] for (k, i), v in per_item.items() if k == kind and not (k == "exp" and i in rent_ids))
    res["sum_sheet"], res["sum_edge"], res["sum_exp"], res["sum_owner"] = s("sheet"), s("edge"), s("exp"), s("owner")
    res["rent_paid"] = rent_pm[0]          # тўланган (плюс билан ёзилгани)
    res["rent_in_xl"] = -rent_pm[1]        # тушган/қайтган (минус билан ёзилгани)
    res["rent_xl"] = -res["rent_paid"]
    res["rent_total"] = tot["rent"] + res["rent_in_xl"] - res["rent_paid"]
    res["sum_wage_usd"] = sum(v[2] for v in per_worker.values())
    res["income"] = tot["in1"] + tot["in2"]
    res["outcome"] = res["sum_wage_usd"] + res["sum_sheet"] + res["sum_edge"] + res["sum_exp"]
    res["profit"] = res["income"] - res["outcome"]
    owners = [i for i in items("owner", ym)]
    wsum = sum(o["share"] for o in owners) or 1
    basis = sum(per_item[("owner", o["id"])][2] for o in owners if o["share"])     # «от» H41: улуш фақат номи бор шериклардан
    res["owners"] = [dict(item=o, taken=per_item[("owner", o["id"])][2], due=basis * o["share"] / wsum) for o in owners]
    first, last = dl[0], max((x for x in dl if x["filled"]), key=lambda x: x["day"], default=dl[0])
    res["start_usd"] = first["start"][0] / first["rate"] + first["start"][1]
    res["last"] = last
    res["end_usd"] = (last["fact"] or last["end"])[0] / last["rate"] + (last["fact"] or last["end"])[1]
    # ойлик камомат («от» M35 = кун охири − факт, ҳаммаси $ да кун курси билан; N35 = кунлик фарқлар йиғиндиси)
    fl = [x for x in dl if x["fact"] is not None]
    res["month_diff"] = res["month_diff_day"] = None
    res["month_diff_parts"] = res["diff_parts"] = (0.0, 0.0)
    if fl:
        lf = fl[-1]
        upto = [x for x in dl if x["day"] <= lf["day"]]
        flow = sum((x["inc"][0] - x["out"][0]) / x["rate"] + x["inc"][1] - x["out"][1] for x in upto)
        res["month_diff"] = first["start"][0] / first["rate"] + first["start"][1] + flow - (lf["fact"][0] / lf["rate"] + lf["fact"][1])
        res["month_diff_day"] = lf["day"]
        res["month_diff_parts"] = (first["start"][0] + sum(x["inc"][0] - x["out"][0] for x in upto) - lf["fact"][0],
                                   first["start"][1] + sum(x["inc"][1] - x["out"][1] for x in upto) - lf["fact"][1])
        res["diff_parts"] = (sum(x["end"][0] - x["fact"][0] for x in fl), sum(x["end"][1] - x["fact"][1] for x in fl))
    _memo[key] = res
    return res

def ensure_kpi_month(ym):
    """Ой учун KPI устунлари: йўқ бўлса — олдинги ойдагидек (ставкалари билан), Excel'да янги ой файли каби."""
    db = A.db()
    if A.q("SELECT 1 FROM k_kpi_mcols WHERE ym=? LIMIT 1", (ym,)):
        return
    prev = A.q("SELECT MAX(ym) m FROM k_kpi_mcols WHERE ym<?", (ym,), one=True)["m"]
    if prev:
        db.execute("INSERT OR IGNORE INTO k_kpi_mcols(ym,col_id,rate,sort,xl_col) SELECT ?, m.col_id, m.rate, m.sort, m.xl_col FROM k_kpi_mcols m "
                   "JOIN k_kpi_cols c ON c.id=m.col_id JOIN k_workers w ON w.id=c.worker_id WHERE m.ym=? AND c.active=1 AND (w.fired IS NULL OR w.fired='' OR w.fired>=?)",
                   (ym, prev, ym + "-01"))
    else:
        db.execute("INSERT OR IGNORE INTO k_kpi_mcols(ym,col_id,rate,sort) SELECT ?, id, rate, sort FROM k_kpi_cols WHERE active=1", (ym,))
    db.commit()

def kpi_cols(ym):
    """Шу ойнинг KPI устунлари (+ шу ойда сони бор, лекин рўйхатда йўқлари)."""
    rows = A.q("SELECT c.id, c.worker_id, c.kind, c.shift, c.active, m.rate, m.sort, m.xl_col, w.name FROM k_kpi_mcols m JOIN k_kpi_cols c ON c.id=m.col_id "
               "JOIN k_workers w ON w.id=c.worker_id WHERE m.ym=? ORDER BY c.kind DESC, m.sort, w.name, c.shift", (ym,))
    have = {r["id"] for r in rows}
    extra = A.q("SELECT DISTINCT c.id, c.worker_id, c.kind, c.shift, c.active, c.rate, c.sort, '' xl_col, w.name FROM k_kpi k JOIN k_kpi_cols c ON c.id=k.col_id "
                "JOIN k_workers w ON w.id=c.worker_id WHERE k.ym=?", (ym,))
    return list(rows) + [r for r in extra if r["id"] not in have]

def kpi_month(ym):
    """KPI: {worker_id: сумма}, {col_id: qty} — ставка шу ойники."""
    cols = {c["id"]: c for c in kpi_cols(ym)}
    by_w, by_c = defaultdict(float), defaultdict(float)
    for r in A.q("SELECT col_id, SUM(qty) q FROM k_kpi WHERE ym=? GROUP BY col_id", (ym,)):
        c = cols.get(r["col_id"])
        if c:
            by_c[c["id"]] = r["q"]
            by_w[c["worker_id"]] += r["q"] * c["rate"]
    return by_w, by_c

def staff_calc(ym, avg, per_worker, _depth=0):
    """Ойлик (как лист «т» + колонки E–K листа «от»)."""
    norm = fnum(A.get_setting("kassa_norm_days", "28")) or 28
    hours = fnum(A.get_setting("kassa_day_hours", "9")) or 9
    kpi_w, _ = kpi_month(ym)
    try:
        import erp as _erp
        rded = _erp.resp_ded(ym)
    except Exception:
        rded = {}
    prev = None
    rows = []
    ws = A.q("SELECT s.*, w.name, w.cur, w.active, w.fired, w.svc FROM k_ws s JOIN k_workers w ON w.id=s.worker_id WHERE s.ym=? ORDER BY w.svc, s.sort, w.name", (ym,))
    listed = {r["worker_id"] for r in ws}
    extra_ids = [wid for wid in per_worker if wid not in listed]
    for wid in extra_ids:   # выдали деньги, а в списке месяца человека нет
        w = A.q("SELECT * FROM k_workers WHERE id=?", (wid,), one=True)
        if w:
            ws.append(dict(worker_id=wid, name=w["name"], cur=w["cur"], active=w["active"], fired=w["fired"], svc=w["svc"], sort=9999, salary=0, days_override=None,
                           extra=0, total_override=None, opening=0, opening_set=0, ts=""))
    for s in ws:
        days, marks = ts_days(s["ts"], hours)
        if s["days_override"] is not None:
            days = s["days_override"]
        base = s["salary"] / norm * days if s["salary"] else 0.0
        kpi = kpi_w.get(s["worker_id"], 0.0)
        ded = rded.get(s["worker_id"], 0.0)                    # масъул шахс сифатида ойликдан ушланган (накладной)
        calc_total = base + kpi + (s["extra"] or 0) - ded
        total = s["total_override"] if s["total_override"] is not None else calc_total
        if s["opening_set"]:
            opening = s["opening"]
        else:
            if prev is None:
                pym = ym_shift(ym, -1)
                prev = {r["worker_id"]: r["closing"] for r in month_calc(pym, _depth + 1)["staff"]} if (_depth < 36 and has_data(pym)) else {}
            opening = prev.get(s["worker_id"], 0.0)
        pu = per_worker.get(s["worker_id"], [0.0, 0.0, 0.0, 0.0])
        paid = pu[0] / avg + pu[1] if s["cur"] == "USD" else pu[0] + pu[3]
        fired = s["fired"] if s["fired"] and s["fired"][:7] == ym else None
        closing = total + opening - paid
        k = avg if s["cur"] == "USD" else 1 / avg      # иккинчи валютага ўгириш
        rows.append(dict(s=s, worker_id=s["worker_id"], name=s["name"], cur=s["cur"], days=days, marks=marks, base=base, kpi=kpi,
                         calc_total=calc_total, total=total, ded=ded, opening=opening, paid=paid, paid_uzs=pu[0], paid_usd_cash=pu[1], paid_usd=pu[2],
                         closing=closing, alt=dict(total=total * k, opening=opening * k, paid=paid * k, closing=closing * k),
                         svc=bool(s["svc"]), fired=fired, fired_day=int(fired[8:10]) if fired else None, override=s["total_override"] is not None))
    return rows

# ---------------------------------------------------------------- нақдсиз ҳисоблар (карта, банк, инвестиция)
def accounts():
    """Нақд кассадан бошқа ҳисоблар (erp_accounts). Жадвал йўқ бўлса — бўш."""
    try:
        return A.q("SELECT * FROM erp_accounts WHERE kind<>'cash' AND active=1 ORDER BY sort, id")
    except Exception:
        return []


def _acc_changes(acc, d1, d2):
    """Ҳисоб бўйича кунлик ўзгаришлар {date: [кирим_сўм, кирим_$, чиқим_сўм, чиқим_$]}:
    1С'даги «Поступление денег» қаторлари (шу ҳисоб «мижоз»ига минус — ҳисобга тушган, плюс — ҳисобдан олинган);
    кассадаги «нақдсиз тўлов»лар (чиқим); ҳисоблар орасидаги ўтказмалар."""
    ch = defaultdict(lambda: [0.0, 0.0, 0.0, 0.0])
    ext = (acc["c1_ext"] if "c1_ext" in acc.keys() else "") or ""
    if ext:
        for r in A.q("SELECT d.date, l.uzs, l.usd FROM erp_lines l JOIN erp_docs d ON d.id=l.doc_id WHERE d.kind='pay' AND d.src IN ('1c','pf') "
                     "AND d.status='posted' AND l.client=? AND d.date BETWEEN ? AND ?", (ext, d1, d2)):
            c = ch[r["date"]]
            for i, v in ((0, r["uzs"] or 0), (1, r["usd"] or 0)):
                if v < 0:
                    c[i] += -v
                else:
                    c[i + 2] += v
    for r in A.q("SELECT ym, day, uzs, usd FROM k_noncash WHERE acc_id=?", (acc["id"],)):
        dt = f"{r['ym']}-{r['day']:02d}"
        if d1 <= dt <= d2:
            ch[dt][2] += r["uzs"]; ch[dt][3] += r["usd"]
    try:
        for r in A.q("SELECT * FROM erp_moves WHERE deleted=0 AND date BETWEEN ? AND ? AND (from_acc=? OR to_acc=?)", (d1, d2, acc["id"], acc["id"])):
            i = 0 if r["cur"] == "uzs" else 1
            if r["to_acc"] == acc["id"]:
                ch[r["date"]][i] += r["amount"]
            if r["from_acc"] == acc["id"]:
                ch[r["date"]][i + 2] += r["amount"]
    except Exception:
        pass
    return ch


def acc_day(acc, d):
    """Ҳисоб кун бошида/охирида: dict(start=[сўм,$], inc=[..], out=[..], end=[..]). Бошланғич қолдиқ — k_acc_open (админ киритади)."""
    ds = d.isoformat()
    op = A.q("SELECT * FROM k_acc_open WHERE acc_id=? AND date<=? ORDER BY date DESC, id DESC LIMIT 1", (acc["id"], ds), one=True)
    base_d = op["date"] if op else "2000-01-01"
    start = [op["uzs"] if op else 0.0, op["usd"] if op else 0.0]
    ch = _acc_changes(acc, base_d, ds)
    for dt, c in ch.items():
        if dt < ds:
            start[0] += c[0] - c[2]; start[1] += c[1] - c[3]
    t = ch.get(ds, [0.0, 0.0, 0.0, 0.0])
    return dict(acc=acc, start=start, inc=[t[0], t[1]], out=[t[2], t[3]], end=[start[0] + t[0] - t[2], start[1] + t[1] - t[3]],
                opened=bool(op))


def day_view(ym, d):
    mc = month_calc(ym)
    return mc, mc["days"][d - 1]


# ---------------------------------------------------------------- Excel import
MONTH_WORDS = {w.lower()[:3]: i for i, w in enumerate(MONTHS) if w}
MONTH_WORDS.update({"янв": 1, "фев": 2, "мар": 3, "апр": 4, "май": 5, "мая": 5, "июн": 6, "июл": 7, "авг": 8, "сен": 9, "окт": 10, "ноя": 11, "дек": 12})

def _month_from_text(t):
    t = str(t or "").lower()
    for w in re.findall(r"[а-яё]+", t):
        if w[:3] in MONTH_WORDS:
            m = MONTH_WORDS[w[:3]]
            y = re.findall(r"20\d\d", t.replace(" ", ""))
            return m, (int(y[0]) if y else None)
    return None, None

def _v(ws, addr):
    v = ws[addr].value
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0

def _s(ws, addr):
    v = ws[addr].value
    if v is None or isinstance(v, (int, float)):
        return ""
    return str(v).strip()

def _expr(wf, addr):
    f = wf[addr].value
    if isinstance(f, str) and f.startswith("=") and _EXPR_OK.match(f[1:]):
        return f[1:]
    return ""

def _note(wf, addr):
    """Excel примечаниеси (муаллиф сатрисиз)."""
    try:
        c = wf[addr].comment
    except Exception:
        return ""
    if not c or not c.text:
        return ""
    t = c.text.strip()
    if c.author and t.startswith(c.author + ":"):
        t = t[len(c.author) + 1:].strip()
    return t[:1000]

def _get_item(db, kind, name, active=1):
    r = db.execute("SELECT id FROM k_items WHERE kind=? AND name=?", (kind, name)).fetchone()
    if r:
        return r[0]
    mx = db.execute("SELECT COALESCE(MAX(sort),0) FROM k_items WHERE kind=?", (kind,)).fetchone()[0]
    return db.execute("INSERT INTO k_items(kind,name,sort,active,grp) VALUES(?,?,?,?,?)",
                      (kind, name, mx + 10, active, exp_group_guess(name) if kind == "exp" else "")).lastrowid

def _lev1(a, b):
    """Отличаются не больше чем одной буквой («Жаср ака» ~ «Жасур ака», «Олмас» ~ «Улмас»)."""
    a, b = a.lower(), b.lower()
    if a == b:
        return True
    if abs(len(a) - len(b)) > 1:
        return False
    if len(a) == len(b):
        return sum(x != y for x, y in zip(a, b)) == 1
    if len(a) > len(b):
        a, b = b, a
    return any(b[:i] + b[i + 1:] == a for i in range(len(b)))

def _match_worker(db, name, month_names):
    """KPI-лист пишет имена по-своему: ищем того же человека среди работников месяца."""
    if name in month_names:
        return _get_worker(db, name)
    near = [n for n in month_names if _lev1(n, name)]
    return _get_worker(db, near[0] if len(near) == 1 else name)

def _get_worker(db, name, cur="UZS"):
    r = db.execute("SELECT id FROM k_workers WHERE name=?", (name,)).fetchone()
    if r:
        if cur == "USD":
            db.execute("UPDATE k_workers SET cur='USD' WHERE id=?", (r[0],))
        return r[0]
    mx = db.execute("SELECT COALESCE(MAX(sort),0) FROM k_workers").fetchone()[0]
    return db.execute("INSERT INTO k_workers(name,cur,sort) VALUES(?,?,?)", (name, cur, mx + 10)).lastrowid

def import_xlsx(path, ym=None, replace=True):
    """Один файл «касса <Месяц> <Год>.xlsx» → месяц в платформе. Берутся сохранённые Excel значения (Excel пересчитывает их при сохранении)."""
    import openpyxl
    wv = openpyxl.load_workbook(path, data_only=True)
    wf = openpyxl.load_workbook(path, data_only=False)
    if "от" not in wv.sheetnames:
        raise ValueError("В файле нет листа «от»")
    ot, otf = wv["от"], wf["от"]
    if not ym:
        m, y = _month_from_text(ot["F1"].value)
        if not m or not y:
            m2, y2 = _month_from_text(os.path.basename(path))
            m, y = m or m2, y or y2
        if not m or not y:
            raise ValueError("Не понял месяц и год (ни в «от»!F1, ни в имени файла)")
        ym = ym_of(y, m)
    lock = (A.get_setting("kassa_platform_from", "") or "").strip()[:7]
    if lock and ym >= lock:
        raise ValueError(f"{ym_title(ym)}: касса платформада юритилади ({lock} дан) — Excel'дан қайта олинмайди")
    db = A.db()
    if replace:
        for t in ("k_days", "k_cells", "k_ws", "k_kpi", "k_meals"):
            db.execute(f"DELETE FROM {t} WHERE ym=?", (ym,))
        db.execute("DELETE FROM k_months WHERE ym=?", (ym,))
    st = defaultdict(int)

    # строки (как на листе «от»)
    sheet_rows = {r: _s(ot, f"B{r}") for r in range(5, 22)}
    edge_rows = {r: _s(ot, f"B{r}") for r in range(25, 33)}
    exp_rows = {r: _s(ot, f"M{r}") for r in range(5, 28)}
    own_rows = {r: _s(ot, f"M{r}") for r in range(29, 32)}
    misc_rows = {33: _s(ot, "M33") or "Аренда", 34: _s(ot, "M34") or "Комуналка"}
    wage_rows, used_names = {}, set()
    for r in list(range(5, 29)) + [30, 31, 32]:
        name = _s(ot, f"E{r}")
        if name in ("Жами:", "Жами ойлик"):
            continue
        if not name:
            # номсиз қатор, лекин кунларда пул берилган — йўқолмасин (кассадан чиққан пул)
            if r < 29 and any(_v(wv[str(d)], f"I{r}") for d in range(1, 32) if str(d) in wv.sheetnames):
                name = f"Номсиз ({r}-қатор)"
            else:
                continue
        base, k2 = name, 2
        while name in used_names:          # два «Дилмурод» в одном месяце — разные люди
            name = f"{base} {k2}"; k2 += 1
        used_names.add(name)
        hf = otf[f"H{r}"].value
        cur = "USD" if isinstance(hf, str) and "/" in hf else "UZS"
        wage_rows[r] = (_get_worker(db, name, cur), name, cur)
        if r >= 30:
            db.execute("UPDATE k_workers SET svc=1 WHERE id=?", (wage_rows[r][0],))
    ids = {}
    for rows, kind in ((sheet_rows, "sheet"), (edge_rows, "edge"), (exp_rows, "exp"), (own_rows, "owner")):
        for r, name in rows.items():
            if name and not name.startswith("Жами"):
                ids[(kind, r)] = _get_item(db, kind, name)
    for r, name in misc_rows.items():
        if name.lower().startswith("аренда"):
            # плюс — аренда, заплаченная хозяину; минус — полученная от арендаторов. Полученную берём из Excel
            # только за дни ДО начала учёта оплат в платформе (потом она приходит из платформы)
            ids[("exp", r)] = _get_item(db, "exp", "Аренда")
            ids[("rentin", r)] = ids[("exp", r)]
            continue
        ids[("exp", r)] = _get_item(db, "exp", name)

    first_pay = db.execute("SELECT MIN(pay_date) FROM payments WHERE is_demo=0").fetchone()[0] or "9999-12-31"
    # дни
    for d in range(1, 32):
        if str(d) not in wv.sheetnames:
            continue
        w, f = wv[str(d)], wf[str(d)]
        try:
            date(*ym_parse(ym), d)
        except ValueError:
            continue
        cells = []
        for r in sheet_rows:
            if ("sheet", r) in ids:
                cells.append(("sheet", ids[("sheet", r)], _v(w, f"C{r}"), _v(w, f"D{r}"), _expr(f, f"C{r}"), _expr(f, f"D{r}")))
        for r in edge_rows:
            if ("edge", r) in ids:
                cells.append(("edge", ids[("edge", r)], _v(w, f"C{r}"), _v(w, f"D{r}"), _expr(f, f"C{r}"), _expr(f, f"D{r}")))
        for r, (wid, _, _) in wage_rows.items():
            cells.append(("wage", wid, _v(w, f"I{r}"), 0.0, _expr(f, f"I{r}"), ""))
        for r in list(exp_rows) + list(misc_rows):
            if ("exp", r) in ids:
                u_, s_ = _v(w, f"N{r}"), _v(w, f"O{r}")
                if r == 33:                     # аренда: выплаты (плюс) отдельно, полученная (минус) — только до платформы
                    iu, is_ = min(u_, 0.0), min(s_, 0.0)
                    if iu or is_:          # Excel'дан олинган кунларда тушган аренда ҳам Excel'дан (платформа тўловлари санаси кассага мос келмайди)
                        cells.append(("exp", ids[("rentin", r)], iu, is_, "", ""))
                    u_, s_ = max(u_, 0.0), max(s_, 0.0)
                cells.append(("exp", ids[("exp", r)], u_, s_, _expr(f, f"N{r}") if u_ else "", _expr(f, f"O{r}") if s_ else ""))
        for r in own_rows:
            if ("owner", r) in ids:
                cells.append(("owner", ids[("owner", r)], _v(w, f"N{r}"), _v(w, f"O{r}"), _expr(f, f"N{r}"), _expr(f, f"O{r}")))
        # Excel'да жами катакка рақам қўлда ёзилган (формула ўрнига) — фарқи кассадан чиққан пул, номсиз қатор бўлиб тушади
        for tot_r, rows_, kind, nm in ((32, own_rows, "owner", "Шерик (номсиз, жами қўлда)"), (28, exp_rows, "exp", "Харажат (номсиз, жами қўлда)")):
            for col in ("N", "O"):
                fv = f[f"{col}{tot_r}"].value
                if isinstance(fv, (int, float)) and not isinstance(fv, bool):
                    diff = fv - sum(_v(w, f"{col}{r}") for r in rows_)
                    if abs(diff) > 1e-9:
                        iid = _get_item(db, kind, nm, active=0)
                        cells.append((kind, iid, diff if col == "N" else 0.0, diff if col == "O" else 0.0, "", ""))
        any_cell = False
        for kind, rid, u, s, eu, es in cells:
            if u or s:
                db.execute("INSERT INTO k_cells(ym,day,kind,ref_id,uzs,usd,expr_uzs,expr_usd) VALUES(?,?,?,?,?,?,?,?) "
                           "ON CONFLICT(ym,day,kind,ref_id) DO UPDATE SET uzs=uzs+excluded.uzs, usd=usd+excluded.usd",
                           (ym, d, kind, rid, u, s, eu, es))
                any_cell = True; st["cells"] += 1
        # Excel катак изоҳлари (примечание) — шу катакка
        for kind, rows_, cu, cs in (("sheet", sheet_rows, "C", "D"), ("edge", edge_rows, "C", "D"), ("exp", list(exp_rows) + list(misc_rows), "N", "O"),
                                    ("owner", own_rows, "N", "O")):
            for r in rows_:
                if (kind, r) not in ids:
                    continue
                nu, ns = _note(f, f"{cu}{r}"), _note(f, f"{cs}{r}")
                if nu or ns:
                    db.execute("INSERT INTO k_cells(ym,day,kind,ref_id,note_uzs,note_usd) VALUES(?,?,?,?,?,?) "
                               "ON CONFLICT(ym,day,kind,ref_id) DO UPDATE SET note_uzs=excluded.note_uzs, note_usd=excluded.note_usd",
                               (ym, d, kind, ids[(kind, r)], nu, ns))
        for r, (wid, _, _) in wage_rows.items():
            nu = _note(f, f"I{r}")
            if nu:
                db.execute("INSERT INTO k_cells(ym,day,kind,ref_id,note_uzs) VALUES(?,?,'wage',?,?) "
                           "ON CONFLICT(ym,day,kind,ref_id) DO UPDATE SET note_uzs=excluded.note_uzs", (ym, d, wid, nu))
        has_fact = w["N42"].value not in (None, "") or w["O42"].value not in (None, "")
        vals = [_v(w, a) for a in ("N37", "O37", "N38", "O38", "N42", "O42")]
        if any_cell or any(vals):
            db.execute("INSERT INTO k_days(ym,day,rate,in1_uzs,in1_usd,in2_uzs,in2_usd,fact_uzs,fact_usd,has_fact,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                       (ym, d, _v(w, "O1") or None, *vals, 1 if has_fact else 0, A.now()))
            st["days"] += 1

    # месяц
    w1 = wv["1"] if "1" in wv.sheetnames else None
    db.execute("INSERT INTO k_months(ym,open_uzs,open_usd,open_set,onec_usd,onec_zav,source) VALUES(?,?,?,?,?,?,?)",
               (ym, _v(w1, "N36") if w1 else 0, _v(w1, "O36") if w1 else 0, 1 if w1 else 0, None, None,                # 1С рақамлари 1С'дан ўзи олинади
                os.path.basename(path)))
    db.execute("UPDATE k_months SET xl_last_day=(SELECT COALESCE(MAX(day),0) FROM k_days WHERE ym=? AND (has_fact=1 OR in1_uzs<>0 OR in1_usd<>0 OR in2_uzs<>0 OR in2_usd<>0)) WHERE ym=?", (ym, ym))

    # табель и ойлик
    t, tf = (wv["т"], wf["т"]) if "т" in wv.sheetnames else (None, None)
    ts_by_name, sal_by_name = {}, {}
    if t:
        def uniq(name, seen):
            base, k2 = name, 2
            while name in seen:
                name = f"{base} {k2}"; k2 += 1
            seen.add(name)
            return name
        seen1, seen2 = set(), set()
        for r in range(5, 19):
            name = _s(t, f"C{r}")
            if name:
                name = uniq(name, seen1)
                codes = []
                for ci in range(4, 35):
                    v = t.cell(row=r, column=ci).value
                    codes.append("" if v is None else (f"{v:g}" if isinstance(v, (int, float)) else str(v).strip()))
                ts_by_name[name] = "|".join(codes)
        for r in range(20, 32):
            name = _s(t, f"K{r}")
            if name:
                name = uniq(name, seen2)
            nf, qf = tf[f"N{r}"].value, tf[f"Q{r}"].value
            if name and isinstance(nf, (int, float)) and isinstance(qf, str) and re.search(r"SUM\(D\d+:AH\d+\)/9", qf):
                sal_by_name[name] = float(nf)
    for i, (r, (wid, name, cur)) in enumerate(sorted(wage_rows.items())):
        total = ot[f"F{r}"].value
        opening = ot[f"G{r}"].value
        sal = sal_by_name.get(name, 0.0)
        db.execute("INSERT OR REPLACE INTO k_ws(ym,worker_id,sort,salary,total_override,opening,opening_set,ts) VALUES(?,?,?,?,?,?,?,?)",
                   (ym, wid, (i + 1) * 10, sal, float(total) if isinstance(total, (int, float)) else 0.0,
                    float(opening) if isinstance(opening, (int, float)) else 0.0, 1, ts_by_name.get(name, "")))
        if sal:
            db.execute("UPDATE k_workers SET salary=? WHERE id=?", (sal, wid))
        st["workers"] += 1

    # KPI — кимга қайси устун ва қанча ставка: «т» листидаги формулалардан (D20='=KPI!C36', F20='=D20*5000')
    if "KPI" in wv.sheetnames:
        st["kpi"] += import_kpi_sheets(db, wv, wf, ym, [n for _, n, _ in wage_rows.values()])

    # овқат: блок с названием месяца в строке 2
    if "х" in wv.sheetnames:
        x, xf = wv["х"], wf["х"]
        m_idx = ym_parse(ym)[1]
        for ci in range(1, x.max_column + 1):
            mname = x.cell(row=2, column=ci).value
            mm, _ = _month_from_text(mname)
            if mm != m_idx:
                continue
            c0 = ci  # «№»
            for rr in range(4, 36):
                dn = x.cell(row=rr, column=c0).value
                try:
                    dn = int(str(dn).strip())
                except (TypeError, ValueError):
                    continue
                if rr == 4 or not 1 <= dn <= ym_days(ym):
                    continue
                b, l, dd = (fnum(_num(x.cell(row=rr, column=c0 + j).value)) for j in (1, 2, 3))
                if b or l or dd:
                    db.execute("INSERT OR REPLACE INTO k_meals(ym,day,b,l,d) VALUES(?,?,?,?,?)", (ym, dn, b, l, dd))
                    st["meals"] += 1
            prices = []
            for j in (1, 2, 3):
                mt = re.search(r"\*([\d.]+)\)?$", str(xf.cell(row=37, column=c0 + j).value or ""))
                prices.append(float(mt.group(1)) if mt else None)
            paid = _num(x.cell(row=38, column=c0 + 1).value)
            db.execute("UPDATE k_months SET meal_b=COALESCE(?,meal_b), meal_l=COALESCE(?,meal_l), meal_d=COALESCE(?,meal_d), meal_paid=? WHERE ym=?",
                       (*prices, fnum(paid), ym))
            break
    # самый свежий месяц задаёт, какие строки показывать в новых днях (как шаблон нового файла)
    latest = db.execute("SELECT MAX(ym) FROM k_months").fetchone()[0]
    if latest == ym:
        keep = set(ids.values())
        for (iid, kind) in db.execute("SELECT id, kind FROM k_items").fetchall():
            db.execute("UPDATE k_items SET active=? WHERE id=?", (1 if iid in keep else 0, iid))
        keepw = {wid for wid, _, _ in wage_rows.values()}
        kpiw = {r[0] for r in db.execute("SELECT DISTINCT c.worker_id FROM k_kpi k JOIN k_kpi_cols c ON c.id=k.col_id WHERE k.ym=?", (ym,))}
        for (wid,) in db.execute("SELECT id FROM k_workers").fetchall():
            db.execute("UPDATE k_workers SET active=? WHERE id=?", (1 if wid in keepw or wid in kpiw else 0, wid))
    db.commit()
    _memo.clear()
    return ym, dict(st)

def kpi_map(wv, wf):
    """«т» листидан KPI жадвали: [(исм, тур, смена, KPI устуни, ставка)]. Исм ва ставка — «т»дан, устун — формуладан."""
    if "т" not in wf.sheetnames or "KPI" not in wv.sheetnames:
        return []
    tf, k = wf["т"], wv["KPI"]
    from openpyxl.utils import column_index_from_string
    edge_from = None                                          # «Кромка» сарлавҳаси турган устундан ўнгга — кромка
    for ci in range(1, k.max_column + 1):
        if str(k.cell(row=2, column=ci).value or "").strip().lower().startswith("кромка"):
            edge_from = ci; break
    out, seen = [], set()
    for qc, rc, nc, block in (("D", "F", "C", "raspil"), ("AB", "AD", "Y", "kromka")):
        name, prev_named = None, False
        for r in range(20, 32):
            nm = tf[f"{nc}{r}"].value
            if isinstance(nm, str) and nm.strip():
                name, cur, prev_named = nm.strip(), nm.strip(), True
            else:
                cur, prev_named = (name if prev_named else None), False
            m = re.match(r"^=KPI!\$?([A-Z]+)\$?36$", str(tf[f"{qc}{r}"].value or "").replace(" ", ""))
            if not (cur and m):
                continue
            col = m.group(1)
            ci = column_index_from_string(col)
            sh = "n" if str(k.cell(row=4, column=ci).value or "").strip().lower().startswith("ноч") else "d"
            kind = ("kromka" if ci >= edge_from else "raspil") if edge_from else block
            rt = re.search(r"\*\s*([\d.]+)\)?$", str(tf[f"{rc}{r}"].value or "").replace(" ", ""))
            key = (cur, kind, sh)
            if key in seen:
                continue
            seen.add(key)
            out.append(dict(name=cur, kind=kind, shift=sh, col=col, ci=ci, rate=float(rt.group(1)) if rt else 0.0))
    return out

def import_kpi_sheets(db, wv, wf, ym, month_names):
    """KPI ойни бошидан: эски сонлар ва устунлар ўчади, «т»даги жадвал бўйича қайта ёзилади."""
    db.execute("DELETE FROM k_kpi WHERE ym=?", (ym,))
    db.execute("DELETE FROM k_kpi_mcols WHERE ym=?", (ym,))
    k = wv["KPI"]
    n = 0
    for x in kpi_map(wv, wf):
        wid = _match_worker(db, x["name"], month_names)
        row = db.execute("SELECT id FROM k_kpi_cols WHERE worker_id=? AND kind=? AND shift=?", (wid, x["kind"], x["shift"])).fetchone()
        cid = row[0] if row else db.execute("INSERT INTO k_kpi_cols(worker_id,kind,shift,rate,sort) VALUES(?,?,?,?,?)",
                                            (wid, x["kind"], x["shift"], x["rate"], x["ci"] * 10)).lastrowid
        db.execute("UPDATE k_kpi_cols SET rate=?, active=1 WHERE id=?", (x["rate"], cid))
        db.execute("INSERT OR REPLACE INTO k_kpi_mcols(ym,col_id,rate,sort,xl_col) VALUES(?,?,?,?,?)", (ym, cid, x["rate"], x["ci"] * 10, x["col"]))
        for d in range(1, ym_days(ym) + 1):
            v = k.cell(row=4 + d, column=x["ci"]).value
            if isinstance(v, (int, float)) and not isinstance(v, bool) and v:
                db.execute("INSERT INTO k_kpi(ym,day,col_id,qty) VALUES(?,?,?,?) ON CONFLICT(ym,day,col_id) DO UPDATE SET qty=excluded.qty",
                           (ym, d, cid, float(v)))
                n += 1
    return n

def import_kpi_file(path, ym=None):
    """Фақат KPI ни Excel'дан қайта олиш (касса, ойлик ва бошқаларга тегмайди)."""
    import openpyxl
    if ym is None:
        m, y = _month_from_text(os.path.basename(path))
        if not m or not y:
            raise ValueError("Файл номидан ой топилмади")
        ym = ym_of(y, m)
    wv = openpyxl.load_workbook(path, data_only=True)
    wf = openpyxl.load_workbook(path)
    lock = (A.get_setting("kassa_platform_from", "") or "").strip()[:7]
    if lock and ym >= lock:
        return ym, 0
    db = A.db()
    names = [r["name"] for r in A.q("SELECT w.name FROM k_ws s JOIN k_workers w ON w.id=s.worker_id WHERE s.ym=?", (ym,))]
    n = import_kpi_sheets(db, wv, wf, ym, names) if "KPI" in wv.sheetnames else 0
    db.commit()
    _memo.clear()
    return ym, n

def _num(v):
    if isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).replace(" ", "").replace(",", "."))
    except (TypeError, ValueError):
        return 0


def find_files(folder, years=None):
    """Файлы «касса <Месяц> <Год>.xlsx» в папке и подпапках «Отчет <Год>»."""
    out = []
    for root, _, files in os.walk(folder):
        for fn in files:
            if fn.startswith("~$") or not fn.lower().endswith(".xlsx") or not fn.lower().startswith("касса"):
                continue
            m, y = _month_from_text(fn)
            if m and y and (not years or y in years):
                out.append((ym_of(y, m), os.path.join(root, fn)))
    return sorted(out)


# ---------------------------------------------------------------- routes
def register(app_module):
    global A
    A = app_module
    app, lr, ar = A.app, A.login_required, A.admin_required
    init_tables()

    @app.before_request
    def _kassa_memo_reset():
        _memo.clear()

    def cur_ym():
        ym = request.values.get("ym") or ym_today()
        if not ym_parse(ym):
            abort(404)
        return ym

    def ctx(ym, **kw):
        return dict(ym=ym, ym_title=ym_title(ym), prev_ym=ym_shift(ym, -1), next_ym=ym_shift(ym, 1), KINDS=KINDS, MONTHS=MONTHS,
                    EXP_GROUPS=EXP_GROUPS, **kw)

    # ---- бир вақтда бир кунни икки киши таҳрирламасин (1С'даги «заблокирован» каби)
    KLOCK_SEC = 120
    _c = A.connect(); _c.execute("CREATE TABLE IF NOT EXISTS k_edit_lock(key TEXT PRIMARY KEY, by TEXT, at REAL)"); _c.commit(); _c.close()

    def _me():
        u = A.current_user()
        return (u["username"] if u else "") or ""

    def klock_get(key):
        """Бошқа одам ҳозир таҳрирлаяптими: (ким, неча сония олдин) ёки None."""
        r = A.q("SELECT * FROM k_edit_lock WHERE key=?", (key,), one=True)
        if r and r["by"] != _me() and time.time() - (r["at"] or 0) < KLOCK_SEC:
            return r["by"], int(time.time() - r["at"])
        return None

    @app.route("/kassa/lock", methods=["POST"])
    @lr
    def kassa_lock():
        key = f"day:{cur_ym()}:{int(A.parse_float(request.form.get('d'), 0))}"
        lk = klock_get(key)
        if lk:
            return jsonify(ok=False, by=lk[0])
        db = A.db()
        if request.form.get("release") == "1":
            db.execute("DELETE FROM k_edit_lock WHERE key=? AND by=?", (key, _me()))
        else:
            db.execute("INSERT OR REPLACE INTO k_edit_lock(key,by,at) VALUES(?,?,?)", (key, _me(), time.time()))
        db.commit()
        return jsonify(ok=True)

    @app.route("/kassa")
    @lr
    def kassa_day():
        ym = cur_ym()
        ensure_month(ym)
        n = ym_days(ym)
        today = date.today()
        dflt = today.day if ym == ym_today() else (n if ym < ym_today() else 1)
        d = min(max(int(A.parse_float(request.args.get("d"), dflt)), 1), n)
        oc1 = sync_1c_in(ym, d)                        # «Кирим 1С» — 1С'дан автомат янгиланади (қўлда ўзгартирилмаган бўлса)
        mc, o = day_view(ym, d)
        lk = klock_get(f"day:{ym}:{d}")
        if lk:      # 1С'дагидек: банд кун умуман очилмайди
            return render_template("kassa_locked.html", **ctx(ym, d=d, mc=mc, lock=lk)), 423
        kclosed = A.is_closed(date(*ym_parse(ym), d))
        if not kclosed:   # кунни очган одам уни банд қилади — бошқалар бир вақтда очолмайди (ёпиқ кун фақат кўрилади)
            db = A.db()
            db.execute("INSERT OR REPLACE INTO k_edit_lock(key,by,at) VALUES(?,?,?)", (f"day:{ym}:{d}", _me(), time.time())); db.commit()
        r = o["row"]
        vals = {(c["kind"], c["ref_id"]): c for c in A.q("SELECT * FROM k_cells WHERE ym=? AND day=?", (ym, d))}
        rate = r["rate"] if r and r["rate"] else default_rate(date(*ym_parse(ym), d))
        groups = [(k, KINDS[k], items(k, ym)) for k in ("sheet", "edge", "exp", "owner")]
        staff = [s for s in mc["staff"] if not (s["fired_day"] and d > s["fired_day"] and ("wage", s["worker_id"]) not in vals)]
        return render_template("kassa_day.html", **ctx(ym, d=d, mc=mc, o=o, r=r, vals=vals, rate=rate, groups=groups, staff=staff,
                                                      onec=oc1 if oc1 is not None else onec_day(date(*ym_parse(ym), d)), kclosed=kclosed,
                                                      accs=[acc_day(a, date(*ym_parse(ym), d)) for a in accounts()], nc=[dict(x) for x in o.get("noncash", [])],
                                                      kclosed_text=A.closed_text(date(*ym_parse(ym), d)) if kclosed else "",
                                                      kver=(A.q("SELECT updated_at FROM k_days WHERE ym=? AND day=?", (ym, d), one=True) or {"updated_at": ""})["updated_at"] or ""))

    @app.route("/kassa/day", methods=["POST"])
    @lr
    def kassa_day_save():
        ym = cur_ym()
        d = int(A.parse_float(request.form.get("d"), 1))
        if not 1 <= d <= ym_days(ym):
            abort(400)
        f = request.form
        db = A.db()
        if A.is_closed(date(*ym_parse(ym), d)):
            flash("Сақланмади: " + A.closed_text(date(*ym_parse(ym), d)), "error")
            return redirect(url_for("kassa_day", ym=ym, d=d))
        lk = klock_get(f"day:{ym}:{d}")
        if lk:
            flash(f"Сақланмади — запись заблокирована: {lk[0]} ҳозир {d:02d}.{ym[5:]} кассасини таҳрирлаяпти.", "error")
            return redirect(url_for("kassa_day", ym=ym, d=d))
        cur = A.q("SELECT updated_at FROM k_days WHERE ym=? AND day=?", (ym, d), one=True)
        if "ver" in f and (f.get("ver") or "") != ((cur["updated_at"] if cur else "") or ""):
            flash(f"Сақланмади: {d:02d}.{ym[5:]} кассасини бошқа одам сиз очгандан кейин ўзгартирди ({cur['updated_at'] if cur else ''}). "
                  "Саҳифани янгиланг ва қайта киритинг.", "error")
            return redirect(url_for("kassa_day", ym=ym, d=d))
        errs = []
        def num_of(name):
            raw = f.get(name, "")
            try:
                return calc(raw), (raw.strip().lstrip("=") if re.search(r"[+\-*/]", raw.strip().lstrip("-")) else "")
            except ValueError:
                errs.append(raw)
                return None, ""
        rate = A.parse_float(f.get("rate"), 0) or default_rate(date(*ym_parse(ym), d))
        head = {}
        for k in ("in1_uzs", "in1_usd", "in2_uzs", "in2_usd", "fact_uzs", "fact_usd"):
            v, _ = num_of(k)
            head[k] = v or 0.0
        has_fact = 1 if (f.get("fact_uzs", "").strip() or f.get("fact_usd", "").strip()) else 0
        cells = []
        for key in f:
            m = re.match(r"^c_(sheet|edge|exp|owner|wage)_(\d+)_(uzs|usd)$", key)
            if m:
                cells.append((m.group(1), int(m.group(2)), m.group(3)))
        if errs:
            flash("Не понял суммы: " + ", ".join(errs[:5]) + ". Пишите числа или выражения вроде 150000+120000", "error")
            return redirect(url_for("kassa_day", ym=ym, d=d))
        oc = onec_day(date(*ym_parse(ym), d))         # «Кирим 1С» 1С'дагидан фарқ қилса — қўлда ўзгартирилган (автомат янгиланмайди)
        man = {}
        for fk in ("in1", "in2"):
            if oc:
                u1, v1 = _in1c(oc, fk)
                man[fk] = 0 if (_same(head[fk + "_uzs"], u1, 0.5) and _same(head[fk + "_usd"], v1, 0.005)) else 1
            else:
                man[fk] = None
        db.execute("INSERT INTO k_days(ym,day,rate,in1_uzs,in1_usd,in2_uzs,in2_usd,fact_uzs,fact_usd,has_fact,note,updated_at,in1_man,in2_man) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?) "
                   "ON CONFLICT(ym,day) DO UPDATE SET rate=excluded.rate, in1_uzs=excluded.in1_uzs, in1_usd=excluded.in1_usd, in2_uzs=excluded.in2_uzs, "
                   "in2_usd=excluded.in2_usd, fact_uzs=excluded.fact_uzs, fact_usd=excluded.fact_usd, has_fact=excluded.has_fact, note=excluded.note, updated_at=excluded.updated_at, "
                   "in1_man=COALESCE(excluded.in1_man, k_days.in1_man), in2_man=COALESCE(excluded.in2_man, k_days.in2_man)",
                   (ym, d, rate, head["in1_uzs"], head["in1_usd"], head["in2_uzs"], head["in2_usd"], head["fact_uzs"], head["fact_usd"], has_fact,
                    (f.get("note") or "").strip()[:500], A.now(), man["in1"], man["in2"]))
        seen = defaultdict(dict)
        notes = defaultdict(dict)
        for kind, rid, cur in cells:
            v, e = num_of(f"c_{kind}_{rid}_{cur}")
            seen[(kind, rid)][cur] = (v or 0.0, e)
            notes[(kind, rid)][cur] = (f.get(f"n_{kind}_{rid}_{cur}") or "").strip()[:1000]
        for (kind, rid), cv in seen.items():
            u, eu = cv.get("uzs", (0.0, ""))
            s, es = cv.get("usd", (0.0, ""))
            nu, ns = notes[(kind, rid)].get("uzs", ""), notes[(kind, rid)].get("usd", "")
            if u or s or nu or ns:
                db.execute("INSERT INTO k_cells(ym,day,kind,ref_id,uzs,usd,expr_uzs,expr_usd,note_uzs,note_usd) VALUES(?,?,?,?,?,?,?,?,?,?) "
                           "ON CONFLICT(ym,day,kind,ref_id) DO UPDATE SET uzs=excluded.uzs, usd=excluded.usd, expr_uzs=excluded.expr_uzs, expr_usd=excluded.expr_usd, "
                           "note_uzs=excluded.note_uzs, note_usd=excluded.note_usd",
                           (ym, d, kind, rid, u, s, eu, es, nu, ns))
            else:
                db.execute("DELETE FROM k_cells WHERE ym=? AND day=? AND kind=? AND ref_id=?", (ym, d, kind, rid))
        if "nc_acc_present" in f:                      # нақдсиз тўловлар (картадан/банкдан)
            db.execute("DELETE FROM k_noncash WHERE ym=? AND day=?", (ym, d))
            accs_ok = {a["id"] for a in accounts()}
            kinds, refs, uz, us, notes = f.getlist("nc_kind"), f.getlist("nc_ref"), f.getlist("nc_uzs"), f.getlist("nc_usd"), f.getlist("nc_note")
            for i, acc in enumerate(f.getlist("nc_acc")):
                try:
                    u = calc(uz[i]) if i < len(uz) and uz[i].strip() else 0.0
                    s_ = calc(us[i]) if i < len(us) and us[i].strip() else 0.0
                except ValueError:
                    continue
                k = kinds[i] if i < len(kinds) else ""
                ref = int(A.parse_float(refs[i] if i < len(refs) else 0, 0)) or None
                if not (u or s_) or not str(acc).isdigit() or int(acc) not in accs_ok or k not in ("sheet", "edge", "exp", "owner", "wage", "other"):
                    continue
                db.execute("INSERT INTO k_noncash(ym,day,acc_id,kind,ref_id,uzs,usd,note) VALUES(?,?,?,?,?,?,?,?)",
                           (ym, d, int(acc), k, ref, u, s_, (notes[i] if i < len(notes) else "")[:200]))
        db.execute("DELETE FROM k_edit_lock WHERE key=? AND by=?", (f"day:{ym}:{d}", _me()))
        db.commit()
        _memo.clear()
        mc, o = day_view(ym, d)
        msg = f"{d:02d}.{ym[5:]} сақланди"
        if o["diff"] is not None and abs(o["diff"]) >= 1:
            msg += f". Фарқ: {A.usd(o['diff'])} — " + ("кассада пул кам" if o["diff"] > 0 else "кассада пул ортиқча (кўп)")
        flash(msg, "ok" if o["diff"] is None or abs(o["diff"]) < 1 else "warn")
        nxt = d + 1 if f.get("next") and d < ym_days(ym) else d
        return redirect(url_for("kassa_day", ym=ym, d=nxt))

    @app.route("/kassa/month")
    @lr
    def kassa_month():
        ym = cur_ym()
        ensure_month(ym)
        sync_1c_month(ym)
        mc = month_calc(ym)
        groups = [(k, KINDS[k], items(k, ym)) for k in ("sheet", "edge", "exp")]
        return render_template("kassa_month.html", **ctx(ym, mc=mc, groups=groups, oc=onec_month(ym, mc)))

    @app.route("/kassa/month/pdf")
    @lr
    def kassa_month_pdf():
        ym = cur_ym()
        ensure_month(ym)
        data = month_pdf(ym)
        resp = A.app.response_class(data, mimetype="application/pdf")
        resp.headers["Content-Disposition"] = f'attachment; filename="kassa_{ym}.pdf"'
        return resp

    @app.route("/kassa/month/save", methods=["POST"])
    @ar
    def kassa_month_save():
        ym = cur_ym()
        ensure_month(ym)
        f = request.form
        onec, zav = f.get("onec_usd", "").strip(), f.get("onec_zav", "").strip()
        A.db().execute("UPDATE k_months SET onec_usd=?, onec_zav=?, note=? WHERE ym=?",
                       (A.parse_float(onec) if onec else None, A.parse_float(zav) if zav else None, (f.get("note") or "")[:1000], ym))
        if f.get("open_set"):
            A.db().execute("UPDATE k_months SET open_set=1, open_uzs=?, open_usd=? WHERE ym=?",
                           (A.parse_float(f.get("open_uzs")), A.parse_float(f.get("open_usd")), ym))
        elif "open_set" in f or f.get("open_auto"):
            A.db().execute("UPDATE k_months SET open_set=0 WHERE ym=?", (ym,))
        A.db().commit()
        flash("Сақланди", "ok")
        return redirect(url_for("kassa_month", ym=ym))

    @app.route("/kassa/staff", methods=["GET", "POST"])
    @lr
    def kassa_staff():
        ym = cur_ym()
        ensure_month(ym)
        n = ym_days(ym)
        if request.method == "POST":
            f, db = request.form, A.db()
            act = f.get("action")
            if act in ("add", "remove", "fire", "unfire") and not A.is_admin():
                abort(403)
            if act == "add":
                name = (f.get("name") or "").strip()
                if name and not f.get("svc") and not (f.get("fio") or "").strip():
                    flash("Ишчининг Ф.И.Ш.ини ёзинг (Фамилия Исм Шарифи)", "error")
                    name = ""
                if name:
                    wid = _get_worker(db, name, "USD" if f.get("cur") == "USD" else "UZS")
                    db.execute("UPDATE k_workers SET active=1, fired=NULL" + (", svc=1" if f.get("svc") else "") + " WHERE id=?", (wid,))
                    if (f.get("fio") or "").strip():
                        db.execute("UPDATE k_workers SET fio=? WHERE id=?", (re.sub(r"\s+", " ", f.get("fio")).strip()[:120], wid))
                    mx = db.execute("SELECT COALESCE(MAX(sort),0) FROM k_ws WHERE ym=?", (ym,)).fetchone()[0]
                    db.execute("INSERT OR IGNORE INTO k_ws(ym,worker_id,sort,salary) VALUES(?,?,?,?)", (ym, wid, mx + 10, A.parse_float(f.get("salary"))))
                    flash(f"{name} қўшилди" + (" (ташқи хизмат)" if f.get("svc") else ""), "ok")
            elif act == "remove":
                wid = int(A.parse_float(f.get("wid")))
                if A.q("SELECT 1 FROM k_cells WHERE ym=? AND kind='wage' AND ref_id=? AND (uzs<>0 OR usd<>0)", (ym, wid)):
                    flash("Бу ойда унга пул берилган — олиб ташлаб бўлмайди", "error")
                else:
                    db.execute("DELETE FROM k_ws WHERE ym=? AND worker_id=?", (ym, wid))
                    flash("Рўйхатдан олинди (фақат шу ой)", "ok")
            elif act == "fire":
                wid = int(A.parse_float(f.get("wid")))
                fd = (f.get("fired") or "").strip()
                try:
                    date.fromisoformat(fd)
                except ValueError:
                    fd = date.today().isoformat() if ym == ym_today() else f"{ym}-{ym_days(ym):02d}"
                if fd[:7] != ym:
                    fd = f"{ym}-{min(max(int(fd[8:10]), 1), ym_days(ym)):02d}"
                db.execute("UPDATE k_workers SET fired=?, active=0 WHERE id=?", (fd, wid))
                # кейинги ойларда пул берилмаган бўлса — рўйхатдан чиқади
                for (nym,) in db.execute("SELECT ym FROM k_ws WHERE worker_id=? AND ym>?", (wid, ym)).fetchall():
                    if not db.execute("SELECT 1 FROM k_cells WHERE ym=? AND kind='wage' AND ref_id=?", (nym, wid)).fetchone():
                        db.execute("DELETE FROM k_ws WHERE ym=? AND worker_id=?", (nym, wid))
                flash(f"Ишдан бўшади: {fd[8:10]}.{fd[5:7]}. Шу ой итогида қолади, кейинги ойга ўтмайди", "ok")
            elif act == "unfire":
                wid = int(A.parse_float(f.get("wid")))
                db.execute("UPDATE k_workers SET fired=NULL, active=1 WHERE id=?", (wid,))
                flash("Бўшатиш бекор қилинди", "ok")
            else:
                for r in A.q("SELECT * FROM k_ws WHERE ym=?", (ym,)):
                    wid = r["worker_id"]
                    if f"sal_{wid}" not in f:
                        continue
                    ts = "|".join((f.get(f"ts_{wid}_{d}") or "").strip()[:4] for d in range(1, 32))
                    tov = f.get(f"tot_{wid}", "").strip()
                    dov = f.get(f"days_{wid}", "").strip()
                    op = f.get(f"open_{wid}", "").strip()
                    try:
                        extra = calc(f.get(f"extra_{wid}"))
                    except ValueError:
                        extra = 0.0
                    db.execute("UPDATE k_ws SET salary=?, extra=?, total_override=?, days_override=?, opening=?, opening_set=?, ts=? WHERE ym=? AND worker_id=?",
                               (A.parse_float(f.get(f"sal_{wid}")), extra, A.parse_float(tov) if tov else None, A.parse_float(dov) if dov else None,
                                A.parse_float(op) if op else 0.0, 1 if op else 0, ts, ym, wid))
                    db.execute("UPDATE k_workers SET salary=? WHERE id=?", (A.parse_float(f.get(f"sal_{wid}")), wid))
                flash("Табель ва ойлик сақланди", "ok")
            db.commit()
            return redirect(url_for("kassa_staff", ym=ym))
        mc = month_calc(ym)
        wd = [date(*ym_parse(ym), d).weekday() for d in range(1, n + 1)]
        return render_template("kassa_staff.html", **ctx(ym, mc=mc, n=n, wd=wd, ts_list=ts_list,
                               norm=A.get_setting("kassa_norm_days", "28"), hours=A.get_setting("kassa_day_hours", "9")))

    @app.route("/kassa/kpi", methods=["GET", "POST"])
    @lr
    def kassa_kpi():
        ym = cur_ym()
        ensure_month(ym)
        ensure_kpi_month(ym)
        n = ym_days(ym)
        db = A.db()
        if request.method == "POST":
            f = request.form
            if f.get("action") in ("addcol", "hide") and not A.is_admin():
                abort(403)
            if f.get("action") == "addcol":
                name = (f.get("name") or "").strip()
                if name and f.get("kind") in KPI_KINDS:
                    wid = _get_worker(db, name)
                    mx = db.execute("SELECT COALESCE(MAX(sort),0) FROM k_kpi_mcols WHERE ym=?", (ym,)).fetchone()[0]
                    for i, sh in enumerate(("d", "n")):
                        rate = A.parse_float(f.get(f"rate_{sh}"))
                        db.execute("INSERT OR IGNORE INTO k_kpi_cols(worker_id,kind,shift,rate) VALUES(?,?,?,?)", (wid, f["kind"], sh, rate))
                        cid = db.execute("SELECT id FROM k_kpi_cols WHERE worker_id=? AND kind=? AND shift=?", (wid, f["kind"], sh)).fetchone()[0]
                        db.execute("UPDATE k_kpi_cols SET active=1 WHERE id=?", (cid,))
                        db.execute("INSERT OR REPLACE INTO k_kpi_mcols(ym,col_id,rate,sort) VALUES(?,?,?,?)", (ym, cid, rate, mx + 10 + i))
                    flash(f"{name} — {KPI_KINDS[f['kind']]} устуни қўшилди", "ok")
            elif f.get("action") == "hide":
                cid = int(A.parse_float(f.get("cid")))
                if A.q("SELECT 1 FROM k_kpi WHERE ym=? AND col_id=? AND qty<>0", (ym, cid)):
                    flash("Бу ойда устунда сонлар бор — аввал уларни ўчиринг", "error")
                else:
                    db.execute("DELETE FROM k_kpi_mcols WHERE ym=? AND col_id=?", (ym, cid))
                    flash("Устун шу ойдан олинди (кейинги ойга ҳам ўтмайди)", "ok")
            elif f.get("action") == "unlock":
                if not A.is_admin():
                    abort(403)
                db.execute("DELETE FROM k_kpi_lock WHERE ym=?", (ym,))
                db.commit()
                for fn in KPI_RESYNC:
                    try:
                        fn(ym)
                    except Exception:
                        A.errlog.exception("kpi resync")
                flash("Қўлда ўзгартирилганлар автоматга қайтарилди", "ok")
            else:
                old = {(r["col_id"], r["day"]): r["qty"] for r in A.q("SELECT * FROM k_kpi WHERE ym=?", (ym,))}
                for c in kpi_cols(ym):
                    if f"rate_{c['id']}" in f:
                        db.execute("INSERT INTO k_kpi_mcols(ym,col_id,rate,sort) VALUES(?,?,?,?) ON CONFLICT(ym,col_id) DO UPDATE SET rate=excluded.rate",
                                   (ym, c["id"], A.parse_float(f.get(f"rate_{c['id']}")), c["sort"]))
                    for d in range(1, n + 1):
                        k = f"q_{c['id']}_{d}"
                        if k in f:
                            try:
                                v = calc(f.get(k))
                            except ValueError:
                                v = 0.0
                            if abs(v - (old.get((c["id"], d)) or 0)) < 0.006:
                                continue                                   # ўзгармаган
                            # қўлда ўзгартирилди — автомат ҳисоб энди бу катакка тегмайди
                            db.execute("INSERT OR REPLACE INTO k_kpi_lock(ym,day,col_id,at) VALUES(?,?,?,?)", (ym, d, c["id"], A.now()))
                            if v:
                                db.execute("INSERT OR REPLACE INTO k_kpi(ym,day,col_id,qty) VALUES(?,?,?,?)", (ym, d, c["id"], v))
                            else:
                                db.execute("DELETE FROM k_kpi WHERE ym=? AND day=? AND col_id=?", (ym, d, c["id"]))
                flash("KPI сақланди", "ok")
            db.commit()
            _memo.clear()
            return redirect(url_for("kassa_kpi", ym=ym))
        cols = kpi_cols(ym)
        qty = {(r["col_id"], r["day"]): r["qty"] for r in A.q("SELECT * FROM k_kpi WHERE ym=?", (ym,))}
        locks = {(r["col_id"], r["day"]) for r in A.q("SELECT col_id, day FROM k_kpi_lock WHERE ym=?", (ym,))}
        by_w, by_c = kpi_month(ym)
        return render_template("kassa_kpi.html", **ctx(ym, n=n, cols=cols, qty=qty, locks=locks, by_c=by_c, by_w=by_w, KPI_KINDS=KPI_KINDS, SHIFTS=SHIFTS))

    @app.route("/kassa/meals", methods=["GET", "POST"])
    @lr
    def kassa_meals():
        ym = cur_ym()
        ensure_month(ym)
        n = ym_days(ym)
        db = A.db()
        if request.method == "POST":
            f = request.form
            for d in range(1, n + 1):
                b, l, dd = (A.parse_float(f.get(f"{k}_{d}")) for k in ("b", "l", "d"))
                if b or l or dd:
                    db.execute("INSERT OR REPLACE INTO k_meals(ym,day,b,l,d) VALUES(?,?,?,?,?)", (ym, d, b, l, dd))
                else:
                    db.execute("DELETE FROM k_meals WHERE ym=? AND day=?", (ym, d))
            try:
                paid = calc(f.get("meal_paid"))
            except ValueError:
                paid = 0.0
            db.execute("UPDATE k_months SET meal_b=?, meal_l=?, meal_d=?, meal_paid=? WHERE ym=?",
                       (A.parse_float(f.get("meal_b"), 12000), A.parse_float(f.get("meal_l"), 17000), A.parse_float(f.get("meal_d"), 17000), paid, ym))
            db.commit()
            flash("Овқат сақланди", "ok")
            return redirect(url_for("kassa_meals", ym=ym))
        meals = {r["day"]: r for r in A.q("SELECT * FROM k_meals WHERE ym=?", (ym,))}
        mr = month_row(ym)
        tb, tl, td = (sum(r[k] for r in meals.values()) for k in ("b", "l", "d"))
        cost = tb * mr["meal_b"] + tl * mr["meal_l"] + td * mr["meal_d"]
        return render_template("kassa_meals.html", **ctx(ym, n=n, meals=meals, mr=mr, tb=tb, tl=tl, td=td, cost=cost))

    @app.route("/kassa/items", methods=["GET", "POST"])
    @ar
    def kassa_items():
        db = A.db()
        if request.method == "POST":
            f = request.form
            act = f.get("action")
            if act == "add":
                kind, name = f.get("kind"), (f.get("name") or "").strip()
                if kind in KINDS and name:
                    iid = _get_item(db, kind, name)
                    db.execute("UPDATE k_items SET active=1, share=? WHERE id=?", (A.parse_float(f.get("share")), iid))
                    if kind == "exp" and f.get("grp") in EXP_GROUP_NAMES:
                        db.execute("UPDATE k_items SET grp=? WHERE id=?", (f.get("grp"), iid))
                    flash(f"«{name}» қўшилди", "ok")
            elif act == "save":
                for r in A.q("SELECT * FROM k_items"):
                    if f"name_{r['id']}" in f:
                        nm = (f.get(f"name_{r['id']}") or "").strip() or r["name"]
                        if A.q("SELECT 1 FROM k_items WHERE kind=? AND name=? AND id<>?", (r["kind"], nm, r["id"])):
                            nm = r["name"]
                        g = f.get(f"grp_{r['id']}", r["grp"])
                        db.execute("UPDATE k_items SET name=?, sort=?, active=?, share=?, grp=? WHERE id=?",
                                   (nm, int(A.parse_float(f.get(f"sort_{r['id']}"), r["sort"])), 1 if f.get(f"act_{r['id']}") else 0,
                                    A.parse_float(f.get(f"share_{r['id']}"), r["share"]), g if g in EXP_GROUP_NAMES else r["grp"], r["id"]))
                for w in A.q("SELECT * FROM k_workers"):
                    if f"wname_{w['id']}" in f:
                        nm = (f.get(f"wname_{w['id']}") or "").strip() or w["name"]
                        if A.q("SELECT 1 FROM k_workers WHERE name=? AND id<>?", (nm, w["id"])):
                            nm = w["name"]
                        db.execute("UPDATE k_workers SET name=?, cur=?, active=?, fio=? WHERE id=?",
                                   (nm, "USD" if f.get(f"wcur_{w['id']}") == "USD" else "UZS", 1 if f.get(f"wact_{w['id']}") else 0,
                                    re.sub(r"\s+", " ", f.get(f"wfio_{w['id']}", w["fio"] or "") or "").strip()[:120], w["id"]))
                if "norm" in f:
                    A.set_settings(commit=False, kassa_norm_days=str(A.parse_float(f.get("norm"), 28) or 28).rstrip("0").rstrip("."),
                                   kassa_day_hours=str(A.parse_float(f.get("hours"), 9) or 9).rstrip("0").rstrip("."))
                A.set_settings(commit=False, kassa_rent="1" if f.get("rent") else "0")
                flash("Сақланди", "ok")
            db.commit()
            return redirect(url_for("kassa_items"))
        return render_template("kassa_items.html", **ctx(ym_today(), all_items=items(all_=True), workers=A.q("SELECT * FROM k_workers ORDER BY active DESC, sort, name"),
                               norm=A.get_setting("kassa_norm_days", "28"), hours=A.get_setting("kassa_day_hours", "9"),
                               rent=A.get_setting("kassa_rent", "1") == "1"))

    @app.route("/kassa/import", methods=["GET", "POST"])
    @ar
    def kassa_import():
        folder = A.get_setting("kassa_folder", r"D:\Отчет")
        if request.method == "POST":
            f = request.form
            folder = (f.get("folder") or folder).strip()
            A.set_settings(kassa_folder=folder)
            years = {int(y) for y in re.findall(r"20\d\d", f.get("years", ""))} or None
            files = find_files(folder, years)
            only = set(f.getlist("ym"))
            if only:
                files = [x for x in files if x[0] in only]
            ok, bad = [], []
            for ym, path in files:
                try:
                    import_xlsx(path, ym)
                    ok.append(ym_title(ym))
                except Exception as e:
                    A.db().rollback()
                    A.errlog.exception("kassa import %s", path)
                    bad.append(f"{os.path.basename(path)}: {e}")
            if ok:
                flash("Кўчирилди: " + ", ".join(ok), "ok")
            if bad:
                flash("Хато: " + "; ".join(bad[:5]), "error")
            if not files:
                flash(f"«{folder}» папкасида «касса …xlsx» файллари топилмади", "error")
            return redirect(url_for("kassa_import"))
        files, err = [], None
        try:
            files = find_files(folder) if os.path.isdir(folder) else []
            if not os.path.isdir(folder):
                err = f"Папка топилмади: {folder}"
        except OSError as e:
            err = str(e)
        have = {r["ym"]: r for r in A.q("SELECT ym, source FROM k_months")}
        return render_template("kassa_import.html", **ctx(ym_today(), folder=folder, files=files, have=have, err=err, lock=A.get_setting("kassa_platform_from", ""),
                                                                   import_result=A.get_setting("kassa_import_result", "")))

    app.jinja_env.globals.update(kassa_ym_title=ym_title)
    _orig_tick = A.scheduler_tick
    def scheduler_tick_kassa(dt=None):
        """«Импортировать при следующем запуске»: сервер сам читает Excel из папки (в т.ч. сетевой \\ADMIN\…)."""
        res = _orig_tick(dt)
        kf = A.get_setting("kassa_kpi_reimport", "")
        if kf:                                       # фақат KPI ни Excel'дан қайта олиш (касса ва ойликка тегмайди)
            A.set_settings(kassa_kpi_reimport="")
            ok, bad = [], []
            try:
                files = find_files(kf)
            except OSError as e:
                files, bad = [], [str(e)]
            for ym, path in files:
                try:
                    _, n_ = import_kpi_file(path, ym); ok.append(f"{ym}:{n_}")
                except Exception as e:
                    A.db().rollback(); bad.append(f"{os.path.basename(path)}: {e}")
                    A.errlog.exception("kassa kpi reimport %s", path)
            A.set_settings(kassa_kpi_result=f"{A.now()}: KPI {', '.join(ok)}" + (f"; хато: {'; '.join(bad[:5])}" if bad else ""))
            A.errlog.warning("kassa KPI reimport: %s", A.get_setting("kassa_kpi_result", ""))
        if A.get_setting("kassa_import_now", "0") == "1":
            A.set_settings(kassa_import_now="0")
            folder = A.get_setting("kassa_import_dir", "") or A.get_setting("kassa_folder", r"D:\Отчет")
            only = {x.strip() for x in A.get_setting("kassa_import_only", "").split(",") if x.strip()}
            A.set_settings(kassa_import_dir="", kassa_import_only="")
            ok, bad = [], []
            try:
                files = [x for x in find_files(folder) if not only or x[0] in only]
            except OSError as e:
                files, bad = [], [str(e)]
            for ym, path in files:
                try:
                    import_xlsx(path, ym); ok.append(ym)
                except Exception as e:
                    A.db().rollback(); bad.append(f"{os.path.basename(path)}: {e}")
                    A.errlog.exception("kassa auto import %s", path)
            A.set_settings(kassa_import_result=f"{A.now()}: кўчирилди {len(ok)} ой ({', '.join(ok)})" + (f"; хато: {'; '.join(bad[:5])}" if bad else "")
                           + ("" if files or bad else f"; «{folder}» да файл топилмади"))
            A.errlog.warning("kassa import: %s", A.get_setting("kassa_import_result", ""))
        return res
    A.scheduler_tick = scheduler_tick_kassa
    app.jinja_env.filters["kin"] = kin


def month_pdf(ym):
    """«Ой итоги» → PDF (тушунарли жадваллар). -> bytes"""
    import pdfgen
    import i18n
    T = i18n.tr                                   # фақат ёзувлар (тизим матни) таржима қилинади, маълумотлар (исмлар) — йўқ
    from datetime import datetime
    usd, som = A.usd, A.som
    mc = month_calc(ym)
    m = mc["month"]
    oc = onec_month(ym, mc)
    neg, pos = "#b3261e", "#3E643C"
    title = T(ym_title(ym).split(" ")[0]) + " " + ym_title(ym).split(" ")[-1]
    doc = pdfgen.Doc(footer=f"Lumberture · {T('Касса — ой итоги')} · {title}")
    # шапка: логотип + сарлавҳа
    logo = None
    ui = A.ui_settings()
    if ui.get("logo"):
        logo = pdfgen.load_image(A.ui_logo_path(ui["logo"]))
    x0 = doc.m
    if logo:
        h = 52
        w = h * logo["w"] / logo["h"]
        doc.image(logo, doc.m, doc.y, w, h)
        x0 = doc.m + w + 12
    doc.text(x0, doc.y + 2, T("Касса — ой итоги"), 18, True, pos)
    doc.text(x0, doc.y + 25, title, 13, True)
    doc.text(x0, doc.y + 42, f"{T('Ўртача курс')} {A.money(round(mc['avg']))} {T('сум')}/$ · {mc['last']['day']} {T('кунгача ёзилган')} · "
             f"{T('тузилди')} {datetime.now().strftime('%d.%m.%Y %H:%M')}", 8, False, "#6b7b76")
    doc.y += 64
    # 4 та асосий рақам
    cards = [(T("Кирим (1С + бошқа)"), usd(mc["income"]), "#1f2a27"), (T("Чиқим (ойлик + завод + харажат)"), usd(mc["outcome"]), "#1f2a27"),
             (T("Қолдиқ (фойда)"), usd(mc["profit"]), pos if mc["profit"] >= 0 else neg),
             (T("Касса: ой боши → охири"), f"{usd(mc['start_usd'])} → {usd(mc['end_usd'])}", "#1f2a27")]
    gap = 8
    cw = (doc.cw - gap * 3) / 4
    for i, (t, v, col) in enumerate(cards):
        x = doc.m + i * (cw + gap)
        doc.rect(x, doc.y, cw, 44, fill="#f0e4d5" if i == 0 else None, stroke="#cfd8d5")
        doc.text(x + 7, doc.y + 7, t, 7.5, False, "#6b7b76")
        sz = 13 if i < 3 else 10.5
        doc.text(x + 7, doc.y + 21, v, sz, True, col)
    doc.y += 58

    def money_rows(kind):
        out = []
        for it in items(kind, ym):
            v = mc["per_item"].get((kind, it["id"]))
            if v and abs(v[2]) >= 0.005:
                out.append((it, v))
        return out

    # 1. Завод
    doc.h1(T("1. Маҳсулотлар учун тўлов (завод), $"))
    rows = []
    for kind, lbl, tot in (("sheet", "Лист", mc["sum_sheet"]), ("edge", "Кромка", mc["sum_edge"])):
        rows.append(dict(cells=[T(lbl + " учун тўлов"), ""], style="sub"))
        rows += [dict(cells=[it["name"], usd(v[2])], indent=10) for it, v in money_rows(kind)]
        rows.append(dict(cells=[T("Жами " + lbl.lower()), usd(tot)], style="tot"))
    rows.append(dict(cells=[T("ЗАВОД ЖАМИ"), usd(mc["sum_sheet"] + mc["sum_edge"])], style="tot"))
    doc.table([(T("Таъминотчи"), 3, "l"), (T("Тўланди, $"), 1, "r")], rows)

    # 2. Харажатлар
    doc.h1(T("2. Харажатлар, $"))
    exp = money_rows("exp")
    rows = []
    for gk, gt in EXP_GROUPS:
        if gk == "rent":
            continue
        g = [(it, v) for it, v in exp if (it["grp"] or "other") == gk]
        if not g:
            continue
        rows.append(dict(cells=[T(gt), usd(sum(v[2] for _, v in g))], style="sub"))
        rows += [dict(cells=[it["name"], usd(v[2])], indent=10) for it, v in g]
    for it, v in exp:
        if (it["grp"] or "other") == "rent" and not it["name"].lower().startswith(T("аренда")):
            rows.append(dict(cells=[it["name"], usd(v[2])], style="sub"))
    rows.append(dict(cells=[T("ЖАМИ ХАРАЖАТ"), usd(mc["sum_exp"])], style="tot"))
    doc.table([(T("Харажат"), 3, "l"), (T("Сумма, $"), 1, "r")], rows)

    # 3. Ойлик
    doc.h1(T("3. Ойлик ва аванс"))

    def two(st, v, alt):
        return f"{usd(v)}\n{som(alt)} {T('сум')}" if st["cur"] == "USD" else f"{som(v)}\n{usd(alt)}"
    rows, svc_hdr = [], False
    for st in mc["staff"]:
        if st["svc"] and not svc_hdr:
            svc_hdr = True
            rows.append(dict(cells=[T("Ташқи хизматлар — бошқа цехларга қилдирилган иш"), "", "", "", "", "", ""], style="sub"))
        nm = st["name"] + (" ($)" if st["cur"] == "USD" else "") + (f" — {T('бўшади')} {st['fired'][8:10]}.{st['fired'][5:7]}" if st["fired"] else "")
        cl = st["closing"]
        rows.append(dict(cells=[nm, two(st, st["total"], st["alt"]["total"]), two(st, st["opening"], st["alt"]["opening"]),
                                som(st["paid_uzs"]) if st["paid_uzs"] else "—", usd(st["paid_usd_cash"]) if st["paid_usd_cash"] else "—",
                                two(st, st["paid"], st["alt"]["paid"]), two(st, cl, st["alt"]["closing"])],
                         color=[None] * 6 + [neg if cl < -0.5 else None], bold_cols=(6,)))
    rows.append(dict(cells=[T("Ойлик жами (берилган пул, кун курси билан)"), "", "", "", "", "", usd(mc["sum_wage_usd"])], style="tot"))
    doc.table([(T("Ишчи"), 2.3, "l"), (T("Ойлик"), 1.1, "r"), (T("Ой бошига"), 1.1, "r"), (T("Олган, сум"), 1.1, "r"), (T("Олган, $"), 0.8, "r"),
               (T("Аванс жами"), 1.1, "r"), (T("Ой охирига"), 1.1, "r")], rows, size=8)
    doc.para(T("Ой охирига: мусбат — ишчига қарзмиз, манфий (қизил) — ортиқча олган. Кичик рақам — иккинчи валютада, ой ўртача курси билан."))

    # 4. Шериклар
    doc.h1(T("4. Шериклар, $"))
    rows = [dict(cells=[o["item"]["name"], usd(o["due"]), usd(o["taken"]), usd(o["due"] - o["taken"])],
                 color=[None, None, None, neg if o["due"] - o["taken"] < -0.5 else None]) for o in mc["owners"]]
    rows.append(dict(cells=[T("Жами"), "", usd(mc["sum_owner"]), ""], style="tot"))
    doc.table([(T("Шерик"), 2, "l"), (T("Улуш"), 1, "r"), (T("Олди"), 1, "r"), (T("Қолди"), 1, "r")], rows)
    doc.para(T("Қолди: манфий — шерик улушидан кўпроқ олган. Ой охирида 1С'да ҳар бирининг ҳисобига ± қилинади."))

    # 5. Кирим ва чиқим + ой якуни
    doc.h1(T("5. Кирим, чиқим ва ой якуни, $"))
    rows = [dict(cells=[T("Касса ой боши"), usd(mc["start_usd"])]),
            dict(cells=[T("Кирим 1С: касса"), usd(mc["tot"]["in1"])], color=[None, pos]),
            dict(cells=[T("Кирим 1С: реализация"), usd(mc["tot"]["in2"])], color=[None, pos]),
            dict(cells=[T("Аренда кирим"), usd(mc["tot"]["rent"] + mc["rent_in_xl"])], color=[None, pos])]
    if mc["rent_paid"]:
        rows.append(dict(cells=[T("Аренда тўланган (кассадан)"), usd(-mc["rent_paid"])]))
    rows += [dict(cells=[T("Чиқим (касса бўйича: шериклар, аренда билан)"), usd(-mc["tot"]["out"])]),
             dict(cells=[f"{T('Касса охири')} ({mc['last']['day']} {T('кун')})", usd(mc["end_usd"])], style="tot"),
             dict(cells=[T("Ой якуни"), ""], style="sub"),
             dict(cells=[T("Ойлик"), usd(mc["sum_wage_usd"])], indent=10),
             dict(cells=[T("Завод (лист + кромка)"), usd(mc["sum_sheet"] + mc["sum_edge"])], indent=10),
             dict(cells=[T("Харажат"), usd(mc["sum_exp"])], indent=10),
             dict(cells=[T("Чиқим жами"), usd(mc["outcome"])], style="tot"),
             dict(cells=[T("Кирим жами"), usd(mc["income"])], style="tot"),
             dict(cells=[T("ҚОЛДИҚ (ФОЙДА) = кирим − чиқим"), usd(mc["profit"])], style="tot",
                  color=[None, pos if mc["profit"] >= 0 else neg])]
    doc.table([("", 3, "l"), ("$", 1, "r")], rows, head=False)

    # 6. 1С билан солиштириш
    doc.h1(T("6. 1С билан солиштириш, $"))
    c_in = m["onec_usd"] if m and m["onec_usd"] is not None else (oc["inc"] if oc else None)
    c_zv = m["onec_zav"] if m and m["onec_zav"] is not None else (oc["zav"] if oc else None)
    k_zv = mc["sum_sheet"] + mc["sum_edge"]
    rows = [dict(cells=[T("Кирим\nкасса + реализация оплатаси"), usd(mc["income"]), usd(c_in) if c_in is not None else "—",
                        usd(mc["income"] - c_in) if c_in is not None else "—"], bold_cols=(3,)),
            dict(cells=[T("Завод\nкасса: лист + кромка; 1С: товар приходи"), usd(k_zv), usd(c_zv) if c_zv is not None else "—",
                        usd(k_zv - c_zv) if c_zv is not None else "—"], bold_cols=(3,))]
    doc.table([("", 2.4, "l"), (T("Касса"), 1, "r"), (T("1С"), 1, "r"), (T("Фарқ"), 1, "r")], rows)
    md = mc["month_diff"]
    rows = [dict(cells=[T("Ой бўйича") + (f" ({mc['month_diff_day']} {T('кун')})" if mc["month_diff_day"] else ""), som(mc["month_diff_parts"][0]),
                        usd(mc["month_diff_parts"][1]), usd(md) if md is not None else "—"],
                 color=[None, None, None, neg if md is not None and md >= 1 else pos if md is not None and md <= -1 else None], bold_cols=(3,)),
            dict(cells=[T("Кунлик фарқлар йиғиндиси"), som(mc["diff_parts"][0]), usd(mc["diff_parts"][1]), usd(mc["tot"]["diff"])],
                 color=[None, None, None, neg if mc["tot"]["diff"] >= 1 else pos if mc["tot"]["diff"] <= -1 else None], bold_cols=(3,))]
    doc.table([(T("Камомат: + кам, − кўп"), 2.4, "l"), (T("сум"), 1, "r"), ("$", 1, "r"), (T("жами, $"), 1, "r")], rows)

    # 7. Кунлар
    doc.new_page()
    doc.h1(T("7. Кунлар бўйича"))
    rows = []
    for x in mc["days"]:
        if not x["filled"]:
            continue
        rows.append(dict(cells=["%02d" % x["day"], A.money(round(x["rate"])), som(x["inc"][0]), usd(x["inc"][1]), som(x["out"][0]), usd(x["out"][1]),
                                som(x["end"][0]), usd(x["end"][1]), som(x["fact"][0]) if x["fact"] else "—", usd(x["fact"][1]) if x["fact"] else "",
                                usd(x["diff"]) if x["diff"] is not None else ""],
                         color=[None] * 10 + [neg if x["diff"] is not None and x["diff"] >= 1 else pos if x["diff"] is not None and x["diff"] <= -1 else "#6b7b76"]))
    if not rows:
        rows.append(dict(cells=[T("Бу ойда ҳали ёзув йўқ")] + [""] * 10, style="note"))
    doc.table([(T("Кун"), 0.55, "l"), (T("Курс"), 0.8, "r"), (T("Кирим, сум"), 1.35, "r"), (T("Кирим, $"), 1, "r"), (T("Чиқим, сум"), 1.35, "r"), (T("Чиқим, $"), 1, "r"),
               (T("Кун охири, сум"), 1.35, "r"), ("$", 1, "r"), (T("Факт, сум"), 1.35, "r"), ("$", 1, "r"), (T("Фарқ"), 0.9, "r")], rows, size=7, zebra=True)
    doc.para(T("Фарқ: + кам (кассада пул етмайди), − кўп (кассада ортиқча пул)."))
    return doc.output()



def _in1c(oc, f):
    """onec_day натижасидан «Кирим 1С» қатори учун (сўм, $)."""
    return (oc.get("uzs", 0), oc.get("usd", 0)) if f == "in1" else (oc.get("s_uzs", 0), oc.get("s_usd", 0))


def _same(a, b, tol):
    return abs((a or 0) - (b or 0)) <= tol


def sync_1c_in(ym, d, oc=None):
    """«Кирим 1С: касса / реализация» — 1С'дан автомат: кун ёзилган бўлса ва кассир бу қаторни қўлда ўзгартирмаган бўлса,
    1С'даги янги суммалар базага ҳам ёзилади (ой итоги ҳам янгиланади). Ёпиқ кун ва келажак кунлари тегилмайди. -> onec_day"""
    dd = date(*ym_parse(ym), d)
    if dd > date.today() or A.is_closed(dd):
        return oc
    db = A.db()
    r = db.execute("SELECT in1_uzs, in1_usd, in2_uzs, in2_usd, in1_man, in2_man FROM k_days WHERE ym=? AND day=?", (ym, d)).fetchone()
    if not r:
        return oc
    oc = oc if oc is not None else onec_day(dd)
    if not oc:
        return oc
    sets, args = [], []
    for f in ("in1", "in2"):
        if r[f + "_man"] != 0:      # 1 — қўлда ўзгартирилган; NULL — эски ёзув (Excel/олдинги версия): ҳеч қачон ўзи устидан ёзилмайди
            continue
        u, v = _in1c(oc, f)
        u, v = round(u or 0, 2), round(v or 0, 2)
        if not (_same(r[f + "_uzs"], u, 0.5) and _same(r[f + "_usd"], v, 0.005)):
            sets.append(f"{f}_uzs=?, {f}_usd=?")
            args += [u, v]
    if sets:
        db.execute(f"UPDATE k_days SET {', '.join(sets)} WHERE ym=? AND day=?", args + [ym, d])
        db.commit()
    return oc


def sync_1c_month(ym):
    for (d,) in A.db().execute("SELECT day FROM k_days WHERE ym=? AND (in1_man=0 OR in2_man=0)", (ym,)).fetchall():
        try:
            sync_1c_in(ym, d)
        except Exception:
            pass

CUTOVER = "2026-10-26"          # 1С ўчирилган кун: шундан бошлаб ҳужжатлар фақат платформада


def onec_month(ym, mc):
    """Ой бўйича: кирим (Поступление денег + накладнойдаги оплата, кун курси билан $ да) ва товар приходи (завод), $.
    1С архиви + платформада проведен қилинган ҳужжатлар (26.10.2026 дан 1С ўчди — янгилари фақат платформада)."""
    try:
        y, m = ym_parse(ym)
        d1, d2 = date(y, m, 1), date(y, m, ym_days(ym))
        live = lambda doc: doc["posted"] and not doc["deleted"]
        rate = {o["day"]: o["rate"] for o in mc["days"]}
        r_ = lambda d: rate.get(d.day) or mc["avg"] or 1
        inc = zav = 0.0
        try:
            import c1
            inc = sum((r["uzs"] or 0) / r_(r["doc"]["date"]) + (r["usd"] or 0) for r in c1.payments(d1, d2) if live(r["doc"]))
            jd = {k: v for k, v in c1.docs(210, d1, d2).items() if live(v)}
            for h in c1.table("DH210"):
                dd = jd.get(h["IDDOC"])
                if dd:
                    inc += (h.get("SP1164") or 0) / r_(dd["date"]) + (h.get("SP1163") or 0)
            zav = sum(r["amount"] or 0 for r in c1.receipts(d1, d2) if live(r["doc"]))
        except Exception:
            if d2.isoformat() < CUTOVER:
                return None
        a, b = d1.isoformat(), d2.isoformat()
        for r in A.q("SELECT x.date, l.usd, l.uzs FROM erp_lines l JOIN erp_docs x ON x.id=l.doc_id WHERE x.src='pf' AND x.kind='pay' "
                     "AND x.status='posted' AND x.date BETWEEN ? AND ?", (a, b)):
            inc += (r["uzs"] or 0) / r_(date.fromisoformat(r["date"])) + (r["usd"] or 0)
        for r in A.q("SELECT date, pay_usd, pay_uzs FROM erp_docs WHERE src='pf' AND kind='sale' AND status='posted' AND date BETWEEN ? AND ?", (a, b)):
            inc += (r["pay_uzs"] or 0) / r_(date.fromisoformat(r["date"])) + (r["pay_usd"] or 0)
        for r in A.q("SELECT usd FROM erp_docs WHERE src='pf' AND kind='receipt' AND status='posted' AND date BETWEEN ? AND ?", (a, b)):
            zav += r["usd"] or 0
        return dict(inc=inc, zav=zav)
    except Exception:
        return None

def onec_day(d):
    """«Кирим 1С» учун 1С'дан кунлик суммалар (Excel'даги каби икки қатор):
    касса — «Поступление денег» (37-қатор), реализация — накладнойдаги «Сумма оплаты» сум/$ (38-қатор)."""
    try:
        live = lambda doc: doc["posted"] and not doc["deleted"]
        try:                                            # 1С архиви (26.10.2026 дан янги ҳужжат йўқ); ўқилмаса — фақат платформа
            import c1
            rows = [r for r in c1.payments(d, d) if live(r["doc"])]
            jd = {k: v for k, v in c1.docs(210, d, d).items() if live(v)}
            heads = [h for h in c1.table("DH210") if h["IDDOC"] in jd]
        except Exception:
            if d.isoformat() < CUTOVER:
                return None
            rows, heads = [], []
        out = dict(usd=sum(r["usd"] or 0 for r in rows), uzs=sum(r["uzs"] or 0 for r in rows), n=len(rows),
                   s_uzs=sum(h.get("SP1164") or 0 for h in heads), s_usd=sum(h.get("SP1163") or 0 for h in heads), s_n=len(heads), pf=0)
        # платформада ёзилган (1С'да йўқ) тўловлар ва накладнойдаги нақд тўловлар ҳам
        ds = d.isoformat()
        for r in A.q("SELECT l.usd, l.uzs FROM erp_lines l JOIN erp_docs x ON x.id=l.doc_id WHERE x.src='pf' AND x.kind='pay' AND x.status='posted' "
                     "AND x.date=?", (ds,)):
            out["usd"] += r["usd"] or 0; out["uzs"] += r["uzs"] or 0; out["n"] += 1; out["pf"] += 1
        for r in A.q("SELECT pay_usd, pay_uzs FROM erp_docs WHERE src='pf' AND kind='sale' AND status='posted' AND date=?", (ds,)):
            out["s_usd"] += r["pay_usd"] or 0; out["s_uzs"] += r["pay_uzs"] or 0; out["s_n"] += 1; out["pf"] += 1
        return out
    except Exception:
        return None
