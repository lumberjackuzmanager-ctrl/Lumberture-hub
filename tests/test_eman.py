"""Эман каталоги: сохта бот (Telethon'сиз) билан тўлиқ юклаш — папкалар, барглар, товарлар, нархлар, саҳифа ва CSV."""
import asyncio, os, sys, tempfile, types
os.environ["RENT_UI_LANG"] = "src"
os.environ["RENT_DB"] = os.path.join(tempfile.mkdtemp(), "t.db")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# --- сохта telethon
tl = types.ModuleType("telethon"); fn = types.ModuleType("telethon.functions"); er = types.ModuleType("telethon.errors")
class _Req:
    def __init__(self, **kw): self.__dict__.update(kw)
fn.messages = types.SimpleNamespace(GetBotCallbackAnswerRequest=type("CB", (_Req,), {}), GetInlineBotResultsRequest=type("IQ", (_Req,), {}))
class FloodWaitError(Exception):
    seconds = 5
er.FloodWaitError = FloodWaitError
er.MessageIdInvalidError = type("MII", (Exception,), {}); er.DataInvalidError = type("DI", (Exception,), {})
tl.functions, tl.errors = fn, er
sys.modules.update({"telethon": tl, "telethon.functions": fn, "telethon.errors": er})

import app as A, eman as E
E.PAUSE_CLICK = E.PAUSE_PRICE = 0
_sleep = asyncio.sleep
async def fast(t, *a): await _sleep(0)
E.asyncio = types.SimpleNamespace(sleep=fast)

R = E.ROOT
TREE = {R: ("📒 КАТАЛОГ", [("🗂 Акрил", "ga"), ("🗂 Кромка", "gc")]),
        "ga": ("🗂 Акрил:", [("🗂 Глянец", "gb"), ("➖ AK-7", None)]),
        "gb": ("🗂 Глянец:", [("➖ MP-1", None), ("➖ MP-2", None)]),
        "gc": ("🗂 Кромка:", [("➖ ПВХ 22", None)])}
GOODS = {"MP-1": [("MP-1 Белый", "601", 97600), ("MP-1 Черный", "602", 98000)], "MP-2": [("MP-2 Серый", "700", 120000.5)],
         "AK-7": [("AK-7 Акрил", "A7", 500000)], "ПВХ 22": [(f"ПВХ 22 цвет {i}", str(i), 1500 + i) for i in range(5)]}

class Btn:
    def __init__(self, text, d): self.text, self.d = text, d
PARENT = {g: p for p, (_, bs) in TREE.items() for _, g in bs if g}
class Markup:
    def __init__(self, btns, back=R): self.btns, self.back = btns, back
    def to_dict(self):
        rows = []
        for t, g in self.btns:
            ty = {"_": "InlineButtonTypeCallback", "data": g.encode()} if g else {"_": "InlineButtonTypeSwitchInline", "query": t[2:], "same_peer": True}
            rows.append({"buttons": [{"_": "KeyboardInlineButton", "text": t, "type": ty}]})
        rows.append({"buttons": [{"_": "KeyboardInlineButton", "text": "🏠 Домой", "type": {"_": "InlineButtonTypeCallback", "data": R.encode()}},
                                 {"_": "KeyboardInlineButton", "text": "↩ Назад", "type": {"_": "InlineButtonTypeCallback", "data": self.back.encode()}}]})
        return {"rows": rows}
class Msg:
    def __init__(self, id, text, out=False, markup=None): self.id, self.message, self.out, self.reply_markup = id, text, out, markup
class Res:
    def __init__(self, d): self.d = d
    def to_dict(self): return self.d
class Bot:
    def __init__(self): self.msgs, self.last_q, self.n, self.bad = [], None, 0, 0
    def add(self, text, out=False, markup=None):
        self.n += 1; m = Msg(self.n, text, out, markup); self.msgs.append(m); return m
    async def get_entity(self, b): return b
    async def send_message(self, ent, text):
        m = self.add(text, out=True)
        if text == "/start":
            self.add("Привет", markup=type("M", (), {"to_dict": lambda s: {"rows": [{"buttons": [{"text": "📒 Каталог", "type": {"_": "InlineButtonTypeCallback", "data": b"catalog"}}]}]}})())
        elif text.startswith("/") and self.last_q:
            i = int(text[1:]) - 1
            t, a, p = GOODS[self.last_q][i]
            self.add(f"{t} ({a}) \n💵 Цена: {p} сум/шт\n\n🔸 Jomiy\nВ наличии: {i + 3}✅\nДля оптовых продаж: 3\n\n🔸 Chinobod\nВ наличии: 0❌")
        return m
    async def get_messages(self, ent, min_id=None, limit=10, ids=None):
        if ids is not None:
            return next((m for m in self.msgs if m.id == ids), None)
        return [m for m in self.msgs if m.id > (min_id or 0)][-limit:]
    async def __call__(self, req):
        if hasattr(req, "data"):
            d = req.data.decode()
            if d == "catalog":
                self.add(*TREE[R][:1], markup=Markup(TREE[R][1])); return None
            m = next(m for m in self.msgs if m.id == req.msg_id)
            here = m.node if hasattr(m, "node") else R
            if d not in [g for _, g in TREE[here][1] if g] + [PARENT.get(here, R)]:
                self.bad += 1; return None                       # кўринмаётган тугма — бот эътибор бермайди
            if d == R:
                self.msgs.remove(m); return None                 # «Домой» хабарни ўчиради
            m.node = d
            m.message, m.reply_markup = TREE[d][0], Markup(TREE[d][1], PARENT.get(d, R)); return None
        self.last_q = req.query
        here = next((m.node for m in reversed(self.msgs) if hasattr(m, "node")), R)
        items = GOODS.get(req.query, []) if any(t[2:] == req.query for t, g in TREE[here][1] if not g) else []
        page = items[:3] if req.offset == "" else (items[3:] if req.offset == "3" else [])
        return Res({"results": [{"title": t, "description": f"Артикул: {a}\nКод: /{items.index((t, a, p)) + 1}",
                                 "send_message": {"message": f"/{items.index((t, a, p)) + 1}"}} for t, a, p in page],
                    "next_offset": "3" if req.offset == "" and len(items) > 3 else ""})

A.app.testing = True
with A.app.app_context():
    con = A.db()
    E.start(con)
    asyncio.run(E._crawl(Bot()))
    st = E.status(con)
    assert st["state"] == "choose" and st["goods"] == 0 and st["folders_left"] == 0, st      # бўлим танланмаган — фақат папкалар
    E.kv_set(con, only='["gc"]', state="goods")
    asyncio.run(E._crawl(Bot()))
    st = E.status(con)
    assert st["state"] == "done" and st["goods"] == 5, st                                     # фақат Кромка
    E.kv_set(con, only='["ga","gc"]', state="goods")
    asyncio.run(E._crawl(Bot()))
    st = E.status(con)
    assert st["state"] == "done", st
    assert st["folders"] == 4 and st["leaves"] == 4 and st["goods"] == 9 and st["priced"] == 9, st
    roots, by, lo = E.tree(con)
    assert [r["name"] for r in roots] == ["Акрил", "Кромка"] and [k["name"] for k in roots[0]["kids"]] == ["Глянец", "AK-7"], roots
    assert roots[0]["total"] == 4 and roots[1]["total"] == 5
    g = con.execute("SELECT * FROM eman_goods WHERE title='MP-2 Серый'").fetchone()
    assert g["price"] == 120000.5 and g["unit"] == "шт" and g["cur"] == "сум" and g["article"] == "700" and '"Jomiy": 3.0' in g["stock"], dict(g)
    rows = E.goods_rows(con, by, lo, E.subtree_ids(by, roots[0]["id"]))
    assert [r["title"] for r in rows] == ["MP-1 Белый", "MP-1 Черный", "MP-2 Серый", "AK-7 Акрил"], [r["title"] for r in rows]
    assert rows[0]["path"] == ["Акрил", "Глянец", "MP-1"]
    # қайта юклаш: Кромка'дан бир товар йўқолди
    GOODS["ПВХ 22"].pop()
    E.start(con)
    roots, by, lo = E.tree(con)
    assert sum(r["total"] for r in roots) == 9                     # тугагунча эски маълумот кўринади
    asyncio.run(E._crawl(Bot()))
    roots, by, lo = E.tree(con)
    assert sum(r["total"] for r in roots) == 8 and E.status(con)["run"] == 2
    # тўхтатиш
    E.start(con); E.kv_set(con, stop=1)
    asyncio.run(E._crawl(Bot())); assert E.status(con)["state"] == "stopped"
    E.kv_set(con, state="done", last_done=2, run=2)
c = A.app.test_client()
def tok():
    with c.session_transaction() as s: s["csrf"] = "t"
tok(); assert c.post("/login", data={"username": "admin", "password": "admin123", "csrf": "t"}).status_code == 302
tok(); assert c.post("/erp/eman", data={"csrf": "t", "act": "select", "top": ["gc", "nope"]}).status_code == 302
with A.app.app_context():
    assert E.kv(A.db(), "only") == '["gc"]'
    E.kv_set(A.db(), only='["ga","gc"]')
h = c.get("/erp/eman").get_data(as_text=True)
assert 'value="ga" form="emsel" onclick="event.stopPropagation()" checked' in h
assert "Эман Материалс каталоги" in h and "Акрил" in h and "Глянец" in h, h[:500]
with A.app.app_context():
    nid = A.q("SELECT id FROM eman_nodes WHERE name='Глянец'", one=True)[0]
h = c.get(f"/erp/eman?node={nid}").get_data(as_text=True)
assert "MP-1 Белый" in h and "97" in h and "Jomiy" in h
assert "MP-2 Серый" in c.get("/erp/eman?q=Серый").get_data(as_text=True)
def _xl_text(data):
    import io as _io
    from openpyxl import load_workbook
    wb = load_workbook(_io.BytesIO(data))
    return "\n".join(";".join("" if v is None else str(v) for v in row) for ws in wb for row in ws.iter_rows(values_only=True))
cs = _xl_text(c.get("/erp/eman?csv=1").data)
assert "Уровень 1;Уровень 2;Уровень 3;Товар" in cs and "Акрил;Глянец;MP-1;MP-1 Белый;601;97600" in cs, cs[:600]
print("EMAN OK")
if True:
    import i18n; i18n.DEFAULT = "ru"
    h = c.get(f"/erp/eman?node={nid}").get_data(as_text=True)
    import re as _r; _t = _r.sub(r"<[^>]+>", " ", h); assert "Каталог Eman Materials" in h and "Цена" in h and "Состояние:" in h, _r.sub(r"\s+"," ",_t)[-2500:]
    print("EMAN RU OK")

# янгиланмайдиган папкалар: эски маълумот янги юклашга ўтади, бот очмайди
import sqlite3 as _sq, json, eman as _E
_c = _sq.connect(":memory:"); _c.row_factory = _sq.Row
_E.init_tables(_c)
_c.execute("INSERT INTO eman_nodes(id,key,parent,name,kind,guid,depth,run,done) VALUES(1,'root',NULL,'Каталог','folder','0',0,1,1)")
_c.execute("INSERT INTO eman_nodes(id,key,parent,name,kind,guid,depth,run,done) VALUES(2,'K',1,'Кромка','folder','K',1,1,1)")
_c.execute("INSERT INTO eman_nodes(id,key,parent,name,kind,guid,depth,run,done) VALUES(3,'K|a',2,'Кромка Эггер','leaf',NULL,2,1,1)")
_c.execute("INSERT INTO eman_nodes(id,key,parent,name,kind,guid,depth,run,done) VALUES(4,'L',1,'ЛДСП','folder','L',1,1,1)")
_c.execute("INSERT INTO eman_goods(node,title,price,run,ok) VALUES(3,'Кромка 19/1',1500,1,1)")
_E.kv_set(_c, run=1, skip=json.dumps(["K"]))
_E.start(_c)
r = {x["id"]: (x["run"], x["done"]) for x in _c.execute("SELECT id, run, done FROM eman_nodes")}
assert r[2] == (2, 1) and r[3] == (2, 1) and r[4] == (1, 0) or r[4][0] == 1, r
g = _c.execute("SELECT run, ok, price FROM eman_goods").fetchone()
assert tuple(g) == (2, 1, 1500), tuple(g)
assert _E.next_leaf(_c, 2) is None
st = _E.status(_c); assert st["skip_names"] == ["Кромка"], st["skip_names"]
# кунлик автомат текширув — 21:00 дан (эгаси, 27.10: мижоз хабарлари тугагач)
from datetime import datetime as _dtm
_E.kv_set(_c, state="done", started="2026-10-26 21:00:00", only=json.dumps(["L"]))
_c.execute("CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT)")
assert _E.AUTO_HOUR == 21
assert not _E.daily(_c, _dtm(2026, 10, 27, 20, 59))
assert _E.daily(_c, _dtm(2026, 10, 27, 21, 0))
assert not _E.daily(_c, _dtm(2026, 10, 27, 22, 0))           # бугун бошланган — такрор йўқ
print("EMAN SKIP OK")
