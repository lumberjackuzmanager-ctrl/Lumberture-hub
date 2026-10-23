"""Тесты доработок: поиск, Excel клиентов, суммы в сумах без копеек, меню.  python tests/test_extras.py"""
import os, sys, io, tempfile
tmp = tempfile.mkdtemp()
os.environ["RENT_DB"] = os.path.join(tmp, "t.db")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import app as A
from openpyxl import load_workbook

A.app.testing = True
c = A.app.test_client()
def csrf():
    with c.session_transaction() as s:
        s.setdefault("csrf", "t"); return s["csrf"]
def post(url, data=None):
    d = dict(data or {}); d["csrf"] = csrf(); return c.post(url, data=d)

# ---- фильтр som: суммы в сумах — целые
assert A.som(826643.3) == "826 643" and A.som(1500000) == "1 500 000" and A.som(-2500.6) == "−2 501" and A.som(None) == "0"
assert A.cur(12345.5, "UZS") == "12 346 сум" and A.cur(12.5, "USD") == "$12.50"
assert A.money(11809.19) == "11 809.19"          # курс по-прежнему с копейками

# ---- вход + демо-данные
assert post("/login", {"username": "admin", "password": "admin123"}).status_code == 302
assert post("/admin/demo/create").status_code == 302

# ---- меню: основные разделы + «Настройки ▾» + поиск
h = c.get("/").get_data(as_text=True)
assert 'class="navmore' in h and "Безопасность" in h and 'action="/search"' in h
assert ".3</b> сум" not in h

# ---- поиск
r = c.get("/search?q=")
assert r.status_code == 200
h = c.get("/search?q=мебель").get_data(as_text=True)
assert "Арендаторы (" in h and "Мягкая мебель" in h
r = c.get("/search?q=4445566")                     # часть телефона клиента Рустам Ахмедов -> сразу в карточку
assert r.status_code == 302 and "/clients/" in r.headers["Location"]
h = c.get("/search?q=нетакогоимени").get_data(as_text=True)
assert "ничего не найдено" in h
h = c.get("/search?q=<script>").get_data(as_text=True)
assert "<script>" not in h.split("<main>")[1].split("</main>")[0]

# ---- клиенты: Excel и суммы без копеек
h = c.get("/clients").get_data(as_text=True)
assert "/clients.xlsx?show=debtors" in h
for show in ("debtors", "active", "all", "xxx"):
    r = c.get(f"/clients.xlsx?show={show}")
    assert r.status_code == 200 and r.mimetype.endswith("spreadsheetml.sheet")
    ws = load_workbook(io.BytesIO(r.data)).active
    assert ws["A4"].value == "Клиент" and ws.cell(row=ws.max_row, column=1).value == "ИТОГО долг"
    names = [ws.cell(row=i, column=1).value for i in range(5, ws.max_row)]
    if show in ("debtors", "xxx"):
        with A.app.app_context():
            bal = A.client_balances()
        assert len(names) == sum(1 for v in bal.values() if v > 0.5)
        assert ws.cell(row=ws.max_row, column=4).value == sum(round(v) for v in bal.values() if v > 0.5)

# ---- без входа — на логин
c.get("/logout")
assert c.get("/search?q=a").status_code == 302 and c.get("/clients.xlsx").status_code == 302
print("EXTRA TESTS PASSED")
