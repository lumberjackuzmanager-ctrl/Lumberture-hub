"""Касса: расчёт дня/месяца, перенос остатков, ойлик, KPI, овқат, импорт Excel (если есть файлы: RENT_KASSA_TEST=папка)."""
import os, sys, tempfile
os.environ["RENT_DB"] = os.path.join(tempfile.mkdtemp(), "t.db")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import app as A, kassa as K

A.app.testing = True
c = A.app.test_client()
with c.session_transaction() as s: s["csrf"] = "t"
assert c.post("/login", data={"username": "admin", "password": "admin123", "csrf": "t"}).status_code == 302
with c.session_transaction() as s: s["csrf"] = "t"

# выражения как в Excel
assert K.calc("150000+120000+250000") == 520000 and K.calc("=6145000-200000-1250000") == 4695000 and K.calc("") == 0
assert K.calc("1 200 000") == 1200000 and K.calc("12,5") == 12.5
for bad in ("__import__('os')", "2**99", "abc"):
    try:
        K.calc(bad); raise AssertionError(bad)
    except ValueError:
        pass

def post(url, data):
    data = dict(data, csrf="t")
    r = c.post(url, data=data)
    assert r.status_code == 302, (url, r.status_code, r.get_data(as_text=True)[:300])
    return r

YM = "2030-03"
for u in ("/kassa", "/kassa/month", "/kassa/staff", "/kassa/kpi", "/kassa/meals", "/kassa/items", "/kassa/import"):
    r = c.get(u + "?ym=" + YM); assert r.status_code == 200, u
with A.app.app_context():
    sup = A.q("SELECT id FROM k_items WHERE kind='sheet' ORDER BY sort LIMIT 1", one=True)["id"]
    exp = A.q("SELECT id FROM k_items WHERE kind='exp' AND name='Обед'", one=True)["id"]
    own = A.q("SELECT id FROM k_items WHERE kind='owner' AND name='Жамшид'", one=True)["id"]

# ишчи + ставка
post("/kassa/staff", {"ym": YM, "action": "add", "name": "Тест Ишчи", "fio": "Тестов  Ишчи Ботирович", "salary": "2800000"})
with A.app.app_context():
    wid = A.q("SELECT id FROM k_workers WHERE name='Тест Ишчи'", one=True)["id"]
    assert A.q("SELECT fio FROM k_workers WHERE id=?", (wid,), one=True)["fio"] == "Тестов Ишчи Ботирович"
post("/kassa/items", {"action": "save", f"wname_{wid}": "Тест Ишчи", f"wfio_{wid}": "Тестов Ишчи Акмалович", f"wact_{wid}": "1", "norm": "28", "hours": "9"})
with A.app.app_context():
    assert A.q("SELECT fio FROM k_workers WHERE id=?", (wid,), one=True)["fio"] == "Тестов Ишчи Акмалович"
h_ = c.get("/staff").get_data(as_text=True); assert "Тестов Ишчи Акмалович" in h_
assert "Сотрудники" in c.get("/kassa/items").get_data(as_text=True)
c.post("/kassa/staff", data={"ym": YM, "action": "add", "name": "ФИШсиз", "csrf": "t"})
with A.app.app_context():
    assert not A.q("SELECT 1 FROM k_workers WHERE name='ФИШсиз'", one=True)
# месяц: открытие кассы вручную
post("/kassa/month/save", {"ym": YM, "open_set": "1", "open_uzs": "1000000", "open_usd": "500"})
# день 1: кирим 1С 2 000 000 сум + $1000, лист 1 180 000 сум (= $100 по курсу 11 800) + $200, обед 270 000 (выражение), аванс 500 000, шерик $50
post("/kassa/day", {"ym": YM, "d": "1", "rate": "11800", "in1_uzs": "2000000", "in1_usd": "1000",
                    f"c_sheet_{sup}_uzs": "1180000", f"c_sheet_{sup}_usd": "200", f"c_exp_{exp}_uzs": "150000+120000", f"n_exp_{exp}_uzs": "Наклейка-150 000\nОбед-120 000", f"n_sheet_{sup}_usd": "",
                    f"c_wage_{wid}_uzs": "500000", f"c_owner_{own}_usd": "50", "fact_uzs": "2230000", "fact_usd": "1150"})
with A.app.app_context():
    K._memo.clear()
    mc = K.month_calc(YM)
    d1 = mc["days"][0]
    assert d1["end"] == [1000000 + 2000000 - 270000 - 500000, 500 + 1000 - 100 - 200 - 50], d1["end"]
    assert abs(d1["diff"]) < 0.01, d1["diff"]
    assert A.q("SELECT expr_uzs FROM k_cells WHERE ym=? AND kind='exp'", (YM,), one=True)["expr_uzs"] == "150000+120000"
    assert A.q("SELECT note_uzs FROM k_cells WHERE ym=? AND kind='exp'", (YM,), one=True)["note_uzs"] == "Наклейка-150 000\nОбед-120 000"
    h_ = c.get(f"/kassa?ym={YM}&d=1").get_data(as_text=True)
    assert 'data-note="Наклейка-150 000\nОбед-120 000"' in h_, "note not rendered"
    assert abs(mc["sum_sheet"] - 300) < 1e-6 and abs(mc["sum_owner"] - 50) < 1e-6
    # день 2: кун боши = факт 1-го дня
    assert mc["days"][1]["start"] == [2230000, 1150]
# неправильная сумма не сохраняет
r = c.post("/kassa/day", data={"ym": YM, "d": "2", "csrf": "t", "in1_uzs": "abc"})
assert r.status_code == 302
with A.app.app_context():
    assert not A.q("SELECT 1 FROM k_days WHERE ym=? AND day=2", (YM,))

# табель: 14 полных дней → ставка/28*14 = 1 400 000; аванс 500 000 → ой охирига 900 000
ts = {f"ts_{wid}_{d}": ("9" if d <= 14 else "х") for d in range(1, 32)}
post("/kassa/staff", dict(ts, ym=YM, **{f"sal_{wid}": "2800000", f"extra_{wid}": "", f"tot_{wid}": "", f"days_{wid}": "", f"open_{wid}": ""}))
with A.app.app_context():
    K._memo.clear()
    st = next(s for s in K.month_calc(YM)["staff"] if s["worker_id"] == wid)
    assert abs(st["days"] - 14) < 1e-9 and abs(st["total"] - 1400000) < 0.01 and abs(st["closing"] - 900000) < 0.01, st
# KPI: 10 шт × 5000 = 50 000 к ойлику
post("/kassa/kpi", {"ym": YM, "action": "addcol", "name": "Тест Ишчи", "kind": "raspil", "rate_d": "5000", "rate_n": "5500"})
with A.app.app_context():
    cid = A.q("SELECT id FROM k_kpi_cols WHERE worker_id=? AND shift='d'", (wid,), one=True)["id"]
post("/kassa/kpi", {"ym": YM, f"q_{cid}_3": "4+6", f"rate_{cid}": "5000"})
with A.app.app_context():
    K._memo.clear()
    st = next(s for s in K.month_calc(YM)["staff"] if s["worker_id"] == wid)
    assert abs(st["kpi"] - 50000) < 0.01 and abs(st["closing"] - 950000) < 0.01, st
    assert A.q("SELECT 1 FROM k_kpi_lock WHERE ym=? AND day=3 AND col_id=?", (YM, cid), one=True)      # қўлда — қулф
    assert not A.q("SELECT 1 FROM k_kpi_lock WHERE ym=? AND day=4", (YM,), one=True)
post("/kassa/kpi", {"ym": YM, f"q_{cid}_3": "10", f"rate_{cid}": "5000"})                     # ўзгармаган — тегмайди
post("/kassa/kpi", {"ym": YM, "action": "unlock"})
with A.app.app_context():
    assert not A.q("SELECT 1 FROM k_kpi_lock WHERE ym=?", (YM,), one=True)
    assert A.q("SELECT qty FROM k_kpi WHERE ym=? AND day=3 AND col_id=?", (YM, cid), one=True)["qty"] == 10
# следующий месяц: остатки переходят сами
NX = K.ym_shift(YM, 1)
assert c.get("/kassa?ym=" + NX).status_code == 200
with A.app.app_context():
    K._memo.clear()
    mn = K.month_calc(NX)
    assert mn["start"] == [2230000, 1150], mn["start"]
    st2 = next(s for s in mn["staff"] if s["worker_id"] == wid)
    assert abs(st2["opening"] - 950000) < 0.01, st2
# аренда из платформы — приход в кассу
with A.app.app_context():
    tid = A.q("SELECT id FROM tenants LIMIT 1", one=True)
    if tid:
        A.db().execute("INSERT INTO payments(tenant_id,pay_date,amount,currency) VALUES(?,?,?,?)", (tid["id"], f"{NX}-05", 300, "USD"))
        A.db().commit(); K._memo.clear()
        assert K.month_calc(NX)["days"][4]["rent"] == [0, 300]
# овқат
post("/kassa/meals", {"ym": YM, "b_1": "3", "l_1": "7", "d_1": "5", "meal_b": "12000", "meal_l": "17000", "meal_d": "17000", "meal_paid": "100000"})
h = c.get("/kassa/meals?ym=" + YM).get_data(as_text=True)
assert "240 000" in h.replace(" ", " ").replace(" ", " "), "овқат жами"
for u in ("/kassa", "/kassa/month", "/kassa/staff", "/kassa/kpi"):
    h = c.get(u + "?ym=" + YM).get_data(as_text=True)
    assert "Тест Ишчи" in h or u == "/kassa/month", u
assert c.get("/kassa?ym=1999-13").status_code == 404

# аванс в $ + ташқи хизмат + бўшатиш
YM2 = K.ym_shift(YM, 3)
post("/kassa/staff", {"ym": YM2, "action": "add", "name": "Доллар Ишчи", "fio": "Долларов Ишчи", "salary": "0", "cur": "USD"})
post("/kassa/staff", {"ym": YM2, "action": "add", "name": "Ровер цех", "svc": "1"})
h_ = c.get("/staff?show=all").get_data(as_text=True).split("<tbody>")[1].split("</tbody>")[0]; assert "Ровер цех" not in h_ and "Доллар Ишчи" in h_
with A.app.app_context():
    w2 = A.q("SELECT id FROM k_workers WHERE name='Доллар Ишчи'", one=True)["id"]
    sv = A.q("SELECT id, svc FROM k_workers WHERE name='Ровер цех'", one=True); assert sv["svc"] == 1
post("/kassa/day", {"ym": YM2, "d": "2", "rate": "12000", f"c_wage_{wid}_usd": "10", f"c_wage_{w2}_uzs": "120000", f"c_wage_{w2}_usd": "5",
                    f"c_wage_{sv['id']}_uzs": "600000"})
with A.app.app_context():
    K._memo.clear(); mc = K.month_calc(YM2)
    st = {s["worker_id"]: s for s in mc["staff"]}
    assert abs(st[wid]["paid"] - 120000) < 0.01 and abs(st[wid]["paid_usd_cash"] - 10) < 1e-9, st[wid]
    assert abs(st[w2]["paid"] - 15) < 0.01 and abs(st[w2]["alt"]["paid"] - 180000) < 0.5, st[w2]
    assert st[sv["id"]]["svc"] and abs(mc["sum_wage_usd"] - (10 + 15 + 50)) < 0.01, mc["sum_wage_usd"]
    assert mc["days"][1]["out"] == [720000, 15], mc["days"][1]["out"]
h = c.get(f"/kassa?ym={YM2}&d=2").get_data(as_text=True)
assert "Ташқи хизматлар" in h and f"c_wage_{wid}_usd" in h
# бўшатиш 10-кун: 11-кундан кўринмайди, ой итогида бор, кейинги ойга ўтмайди
post("/kassa/staff", {"ym": YM2, "action": "fire", "wid": str(w2), "fired": f"{YM2}-10"})
assert f"c_wage_{w2}_uzs" in c.get(f"/kassa?ym={YM2}&d=10").get_data(as_text=True)
assert f"c_wage_{w2}_uzs" not in c.get(f"/kassa?ym={YM2}&d=11").get_data(as_text=True)
assert "Доллар Ишчи" in c.get(f"/kassa/month?ym={YM2}").get_data(as_text=True)
_pdf = c.get(f"/kassa/month/pdf?ym={YM2}")
assert _pdf.status_code == 200 and _pdf.data[:5] == b"%PDF-" and _pdf.mimetype == "application/pdf"
NX2 = K.ym_shift(YM2, 1)
assert "Доллар Ишчи" not in c.get(f"/kassa/month?ym={NX2}").get_data(as_text=True)
assert "Тест Ишчи" in c.get(f"/kassa/month?ym={NX2}").get_data(as_text=True)
post("/kassa/staff", {"ym": YM2, "action": "unfire", "wid": str(w2)})
with A.app.app_context():
    assert A.q("SELECT fired FROM k_workers WHERE id=?", (w2,), one=True)["fired"] is None

# импорт Excel (если указана папка с файлами)
src = os.environ.get("RENT_KASSA_TEST")
if src and os.path.isdir(src):
    import openpyxl
    with A.app.test_request_context():
        files = K.find_files(src)
        assert files
        for ym, path in files:
            K.import_xlsx(path, ym)
            K._memo.clear()
            mc = K.month_calc(ym)
            wv = openpyxl.load_workbook(path, data_only=True)
            ot = wv["от"]
            assert abs(mc["sum_sheet"] - (ot["C22"].value or 0)) < 0.05, ym
            assert abs(mc["tot"]["in1"] - (ot["N37"].value or 0)) < 0.05, ym
            if ot["M35"].value is not None and mc["month_diff"] is not None and ym == "2026-09":
                assert abs(mc["tot"]["diff"] - ot["N35"].value) < 0.05, (ym, mc["tot"]["diff"], ot["N35"].value)
            # KPI: ҳар устун = «т» даги ставка, сони = KPI!…36 (ишчи номи «т» листидан)
            wf_ = openpyxl.load_workbook(path)
            kp = K.kpi_map(wv, wf_)
            cols = {(kc["name"], kc["kind"], kc["shift"]): kc for kc in K.kpi_cols(ym)}
            _, by_c = K.kpi_month(ym)
            for x in kp:
                kc = next((kc for (nm, kd, sh), kc in cols.items() if kd == x["kind"] and sh == x["shift"] and (nm == x["name"] or K._lev1(nm, x["name"]))), None)
                assert kc is not None, (ym, x)
                assert abs(kc["rate"] - x["rate"]) < 1e-9, (ym, x, kc["rate"])
                assert abs(by_c.get(kc["id"], 0) - (wv["KPI"].cell(row=36, column=x["ci"]).value or 0)) < 1e-6, (ym, x)
            for o in mc["days"]:
                w = wv[str(o["day"])] if str(o["day"]) in wv.sheetnames else None
                stale = w is not None and abs((w["N41"].value or 0) - ((w["N36"].value or 0) + (w["N37"].value or 0) + (w["N38"].value or 0) - (w["N40"].value or 0))) > 1
                if w is not None and not stale and o["row"] is not None and o["row"]["has_fact"] and not (w["N33"].value or w["O33"].value):
                    assert abs(o["end"][0] - (w["N41"].value or 0)) < 1 and abs(o["end"][1] - (w["O41"].value or 0)) < 0.05, (ym, o["day"])
    for u in ("/kassa", "/kassa/month", "/kassa/staff", "/kassa/kpi", "/kassa/meals"):
        assert c.get(u + "?ym=" + files[-1][0]).status_code == 200
    print("  excel import checked:", ", ".join(ym for ym, _ in files))
print("KASSA TESTS PASSED")

# роль «kassa»: фақат касса саҳифалари
with A.app.app_context():
    from werkzeug.security import generate_password_hash
    A.db().execute("INSERT INTO users(username,password_hash,role) VALUES(?,?,?)", ("kassir1", generate_password_hash("Kassa12345"), "kassa")); A.db().commit()
c2 = A.app.test_client()
with c2.session_transaction() as s: s["csrf"] = "t"
assert c2.post("/login", data={"username": "kassir1", "password": "Kassa12345", "csrf": "t"}).status_code == 302
with A.app.app_context():
    A.db().execute("DELETE FROM k_edit_lock"); A.db().commit()   # admin кассадан чиқди
with c2.session_transaction() as s: s["csrf"] = "t"
for u in ("/kassa", "/kassa/month", "/kassa/meals"):
    r = c2.get(u); assert r.status_code == 200, u
    h_ = r.get_data(as_text=True); assert "Итоги месяца" in h_ and "Арендаторы" not in h_ and "kassa/month/save" not in h_, u
for u in ("/", "/tenants", "/1c/clients", "/admin/users", "/kassa/items", "/kassa/staff", "/kassa/kpi"):
    r = c2.get(u); assert r.status_code == 302 and r.headers["Location"].endswith("/kassa"), (u, r.status_code)
assert c2.post("/kassa/staff", data={"ym": YM, "action": "add", "name": "X", "csrf": "t"}).status_code == 403
assert c2.post("/kassa/month/save", data={"ym": YM, "csrf": "t"}).status_code == 403
print("KASSA ROLE OK")

# платформада юритилган ой Excel'дан қайта олинмайди
src = os.environ.get("RENT_KASSA_TEST")
if src and os.path.isdir(src):
    with A.app.test_request_context():
        A.set_settings(kassa_platform_from="2026-10-14")
        oct_ = [p for y, p in K.find_files(src) if y == "2026-10"]
        if oct_:
            try:
                K.import_xlsx(oct_[0], "2026-10"); raise AssertionError("locked month imported")
            except ValueError as e:
                assert "платформада" in str(e)
        A.set_settings(kassa_platform_from="")
    print("LOCK OK")

# харажат бўлимлари ва битта «Аренда»
with A.app.app_context():
    names = {r["name"]: r["grp"] for r in A.q("SELECT name, grp FROM k_items WHERE kind='exp'")}
    assert names.get("Ровер") == "machine" and names.get("Бухгалтер/Налог") == "tax" and names.get("Базис") == "basis", names
    assert names.get("Компьютер ва канцтовар") == "office" and names.get("Доставка") == "delivery"
    assert "Аренда (тушган, Excel)" not in names
h = c.get("/kassa?ym=" + YM + "&d=3").get_data(as_text=True)
assert "Станоклар" in h and "Бухгалтер / налог" in h and "Компьютер ва канцтовар" in h
print("GROUPS OK")

import time as _t0
# ---- бир кунни бир вақтда икки киши очолмасин (очишнинг ўзи банд қилади)
with A.app.app_context():
    A.db().execute("DELETE FROM k_edit_lock"); A.db().commit()
assert c.get(f"/kassa?ym={YM}&d=7").status_code == 200
with A.app.app_context():
    assert A.q("SELECT by FROM k_edit_lock WHERE key=?", (f"day:{YM}:7",), one=True)[0] == "admin"
assert c.get(f"/kassa?ym={YM}&d=7").status_code == 200           # ўзи қайта очса бўлади
with A.app.app_context():
    A.db().execute("UPDATE k_edit_lock SET by='Жасур', at=? WHERE key=?", (_t0.time(), f"day:{YM}:7")); A.db().commit()
assert c.get(f"/kassa?ym={YM}&d=7").status_code == 423
import time as _t
with c.session_transaction() as s: s["csrf"] = "t"
r = c.post(f"/kassa/lock?ym={YM}", data={"d": "5", "csrf": "t"}); assert r.get_json()["ok"] is True
with A.app.app_context():
    A.db().execute("UPDATE k_edit_lock SET by='Жасур', at=? WHERE key=?", (_t.time(), f"day:{YM}:5")); A.db().commit()
r = c.get(f"/kassa?ym={YM}&d=5"); h = r.get_data(as_text=True)
assert r.status_code == 423 and "Запись заблокирована" in h and "<b>Жасур</b> таҳрирлаяпти" in h and 'id="kform"' not in h
assert c.get(f"/kassa?ym={YM}&d=6").status_code == 200       # бошқа кун очилади
with c.session_transaction() as s: s["csrf"] = "t"
assert c.post(f"/kassa/lock?ym={YM}", data={"d": "5", "csrf": "t"}).get_json() == {"ok": False, "by": "Жасур"}
with c.session_transaction() as s: s["csrf"] = "t"
r = c.post("/kassa/day", data={"ym": YM, "d": "5", "csrf": "t", "in1_uzs": "777", "ver": ""}, follow_redirects=True)
assert "запись заблокирована" in r.get_data(as_text=True)
with A.app.app_context():
    assert not A.q("SELECT 1 FROM k_days WHERE ym=? AND day=5 AND in1_uzs=777", (YM,))
    A.db().execute("UPDATE k_edit_lock SET at=? WHERE key=?", (_t.time() - 500, f"day:{YM}:5")); A.db().commit()   # эски қулф ўтиб кетади
assert c.get(f"/kassa?ym={YM}&d=5").status_code == 200
# версия: бошқа одам сиз очгандан кейин сақлаган бўлса — рад
with A.app.app_context():
    A.db().execute("INSERT OR REPLACE INTO k_days(ym,day,rate,updated_at) VALUES(?,5,12000,'2099-01-01 00:00:00')", (YM,)); A.db().commit()
with c.session_transaction() as s: s["csrf"] = "t"
r = c.post("/kassa/day", data={"ym": YM, "d": "5", "csrf": "t", "in1_uzs": "777", "ver": "эски"}, follow_redirects=True)
assert "бошқа одам" in r.get_data(as_text=True)
with c.session_transaction() as s: s["csrf"] = "t"
r = c.post("/kassa/day", data={"ym": YM, "d": "5", "csrf": "t", "in1_uzs": "777", "ver": "2099-01-01 00:00:00"}, follow_redirects=True)
assert "сақланди" in r.get_data(as_text=True)
with A.app.app_context():
    assert A.q("SELECT 1 FROM k_days WHERE ym=? AND day=5 AND in1_uzs=777", (YM,))
    assert A.q("SELECT 1 FROM k_edit_lock WHERE key=? AND by='admin'", (f"day:{YM}:5",))   # сақлагач кун қайта очилди — яна унда
print("KASSA EDIT LOCK OK")

# ---- ёпиқ кун: кассани ўзгартириб бўлмайди
from datetime import date as _d, timedelta as _td
with A.app.app_context():
    A.set_settings(edit_lock_days="4")
od = _d.today() - _td(days=10)
oym = od.strftime("%Y-%m")
h = c.get(f"/kassa?ym={oym}&d={od.day}").get_data(as_text=True); assert "ёпилган" in h and 'data-closed="1"' in h
with c.session_transaction() as s: s["csrf"] = "t"
r = c.post("/kassa/day", data={"ym": oym, "d": str(od.day), "csrf": "t", "in1_uzs": "999"}, follow_redirects=True)
assert "Сақланмади" in r.get_data(as_text=True)
with A.app.app_context():
    assert not A.q("SELECT 1 FROM k_days WHERE ym=? AND day=? AND in1_uzs=999", (oym, od.day))
    assert not A.q("SELECT 1 FROM k_edit_lock WHERE key=?", (f"day:{oym}:{od.day}",))
    A.set_settings(edit_lock_days="0")
print("KASSA CLOSED OK")

# ---- картадан/банкдан тўловлар: ой жамига киради, нақд кассага эмас; ҳисоб қолдиғи
import erp as _erp
from datetime import date
with A.app.app_context():
    A.set_settings(edit_lock_days="0")
    accs = {a["name"]: a for a in K.accounts()}
    if accs:
        aid = accs["Карта АСАКА 05-76"]["id"]
        it = A.q("SELECT id FROM k_items WHERE kind='sheet' AND active=1 LIMIT 1", one=True)[0]
        A.db().execute("INSERT INTO k_acc_open(acc_id,date,uzs,usd) VALUES(?,?,?,?)", (aid, f"{YM}-01", 5000000, 0)); A.db().commit()
        K._memo.clear()
        before = K.day_view(YM, 9)[1]["end"]
if accs:
    with c.session_transaction() as s: s["csrf"] = "t"
    r = c.post("/kassa/day", data={"ym": YM, "d": "9", "csrf": "t", "nc_acc_present": "", "nc_acc": [str(aid)], "nc_kind": ["sheet"],
                                   "nc_ref": [str(it)], "nc_uzs": ["1200000"], "nc_usd": [""], "nc_note": [""]})
    assert r.status_code == 302
    with A.app.app_context():
        K._memo.clear()
        mc, o = K.day_view(YM, 9)
        assert o["end"] == before, (o["end"], before)                       # нақд касса ўзгармади
        assert abs(mc["per_item"][("sheet", it)][0] - 1200000) < 0.01 or mc["per_item"][("sheet", it)][0] >= 1200000
        ad = K.acc_day(accs["Карта АСАКА 05-76"], date(int(YM[:4]), int(YM[5:]), 9))
        assert ad["out"][0] == 1200000 and ad["end"][0] == ad["start"][0] - 1200000 and ad["start"][0] == 5000000
    h = c.get(f"/kassa?ym={YM}&d=9").get_data(as_text=True); assert "Картадан / банкдан тўловлар" in h and "1200000" in h
    assert 'class=hasnc' in h and "картадан: 1 200 000 сум" in h, h[h.find("hasnc")-200:h.find("hasnc")+400]   # юқоридаги ўша қаторда ҳам (26.10)
    # «Нақдга (кассага)»: картадан $100 ечиб нақд олинди — картадан чиқди, нақд кассага кирим
    with c.session_transaction() as s: s["csrf"] = "t"
    r = c.post("/kassa/day", data={"ym": YM, "d": "9", "csrf": "t", "nc_acc_present": "", "nc_acc": [str(aid), str(aid)], "nc_kind": ["sheet", "tocash"],
                                   "nc_ref": [str(it), "0"], "nc_uzs": ["1200000", ""], "nc_usd": ["", "100"], "nc_note": ["", ""]})
    assert r.status_code == 302
    with A.app.app_context():
        K._memo.clear()
        mc, o = K.day_view(YM, 9)
        assert abs(o["end"][1] - (before[1] + 100)) < 0.001 and o["end"][0] == before[0], (o["end"], before)
        ad = K.acc_day(accs["Карта АСАКА 05-76"], date(int(YM[:4]), int(YM[5:]), 9))
        assert ad["out"][1] == 100 and ad["out"][0] == 1200000
    h = c.get(f"/kassa?ym={YM}&d=9").get_data(as_text=True); assert "Ўтказма" in h and '"tocash"' in h
    with c.session_transaction() as s: s["csrf"] = "t"
    c.post("/kassa/day", data={"ym": YM, "d": "9", "csrf": "t", "nc_acc_present": ""})                    # ҳаммаси ўчирилди
    with A.app.app_context():
        assert not A.q("SELECT 1 FROM k_noncash WHERE ym=? AND day=9", (YM,))
# олдиндан белгиланган оклад: фақат шу ой очилганда (кейин қўлда ўзгартирилгани кейинги ойга ўтади)
with A.app.app_context():
    db = A.db()
    w = db.execute("SELECT worker_id, salary FROM k_ws WHERE ym=? LIMIT 1", (YM,)).fetchone()
    nx = K.ym_shift(YM, 1)
    db.execute("DELETE FROM k_ws WHERE ym>=?", (nx,)); db.commit()
    A.set_settings(salary_plan='[{"worker": %d, "from": "%s", "salary": 2000000}]' % (w[0], nx))
    K.ensure_month(nx)
    assert db.execute("SELECT salary FROM k_ws WHERE ym=? AND worker_id=?", (nx, w[0])).fetchone()[0] == 2000000
    db.execute("UPDATE k_ws SET salary=2500000 WHERE ym=? AND worker_id=?", (nx, w[0])); db.commit()
    K.ensure_month(K.ym_shift(nx, 1))
    assert db.execute("SELECT salary FROM k_ws WHERE ym=? AND worker_id=?", (K.ym_shift(nx, 1), w[0])).fetchone()[0] == 2500000
    db.execute("DELETE FROM k_ws WHERE ym>=?", (nx,)); db.execute("DELETE FROM k_months WHERE ym>=?", (nx,)); db.commit()
    A.set_settings(salary_plan="")
print("KASSA NONCASH OK")
