"""4-қадам: ҳисоботлар — Остаток (кратко/подробно/только), Отчет по накладным, Отчет по приходу, Взаиморасчеты."""
import os, sys, tempfile
os.environ["RENT_DB"] = os.path.join(tempfile.mkdtemp(), "t.db"); os.environ["RENT_UI_LANG"] = "src"
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import app as A, erp, erp_reports as R
from datetime import date
A.app.testing = True
erp.stock_now = lambda: {"G1": 5.0, "G2": 0.0}
R._c1_stock_moves = lambda: None
R._c1_moves = lambda: ({}, {"C1": -100.0})          # 1С: C1 бизга $100 қарз
c = A.app.test_client()
def tok():
    with c.session_transaction() as s: s["csrf"] = "t"
tok(); c.post("/login", data={"username": "admin", "password": "admin123", "csrf": "t"})
with A.app.app_context():
    con = A.db()
    con.execute("INSERT INTO erp_goods(ext,code,name,unit,cur,price,parent,folder,deleted,lc,src) VALUES('F1','F1','Плиты','',0,0,'0',1,0,'плиты','1c')")
    for e, n, p in (("G1", "ЛДСП Белый", "F1"), ("G2", "Кромка 19/1", "0")):
        con.execute("INSERT INTO erp_goods(ext,code,name,unit,cur,price,price_in,parent,folder,deleted,lc,src) VALUES(?,?,?,'шт',2,40,30,?,0,0,?,'1c')",
                    (e, e, n, p, erp.lc(n)))
    for e, n, k in (("S1", "Эман база", "supplier"), ("C1", "Мижоз Бир", "client"), ("C2", "Мижоз Икки", "client")):
        con.execute("INSERT INTO erp_clients(ext,code,name,full,phones,parent,folder,deleted,lc,src,ckind) VALUES(?,?,?,'','','0',0,0,?,'1c',?)",
                    (e, e, n, erp.lc(n), k))
    def doc(kind, no, d, cl, usd, lines, src="1c", paid=0, nal=1):
        i = con.execute("INSERT INTO erp_docs(kind,no,date,time,client,usd,uzs,nal,status,src,paid) VALUES(?,?,?,'10:00',?,?,0,?,'posted',?,?)",
                        (kind, no, d, cl, usd, nal, src, paid)).lastrowid
        for n, (g, q, u) in enumerate(lines):
            con.execute("INSERT INTO erp_lines(doc_id,lineno,good,client,qty,price,usd,uzs,debt) VALUES(?,?,?,?,?,?,?,0,?)",
                        (i, n + 1, g, cl if kind == "pay" else "", q, (u / q) if q else 0, u, u if kind == "pay" else 0))
        return i
    doc("receipt", "R1", "2026-10-02", "S1", 400, [("G1", 10, 400)])
    s1 = doc("sale", "101", "2026-10-05", "C1", 120, [("G1", 3, 120)])
    doc("sale", "102", "2026-10-07", "C2", 80, [("G1", 2, 80)], paid=80)
    doc("sale", "103", "2026-10-20", "C2", 40, [("G1", 1, 40)])           # давр (d2=15.10) дан кейин
    doc("pay", "P1", "2026-10-10", "C1", 50, [("", 0, 50)], src="pf")
    con.execute("UPDATE erp_clients SET debt1c=100 WHERE ext='C1'")
    con.commit()
    # қолдиқ: ҳозир 5 (+0 pf). 15.10 даги = 5 + 1 (20.10 даги сотув) = 6; бош = 6 - 10 + 5 = 1
    goods, res = R.stock_report(con, date(2026, 10, 1), date(2026, 10, 15))
    r = res["G1"]
    assert (r["beg"], r["inc"], r["out"], r["end"]) == (1.0, 10.0, 5.0, 6.0), r
    assert [m["bal"] for m in r["moves"]] == [11.0, 8.0, 6.0]
    m = R.mutual(con, date(2026, 10, 1), date(2026, 10, 15))
    x = {r["client"]: r for r in m}
    # C1: ҳозир 100 (1С) − 50 (pf тўлов) = 50 — 1С ҳаракатлари мок қилинмаган, шунинг учун бош = 50 + 50 = 100
    assert x["C1"]["end"] == 50 and x["C1"]["paid"] == 50 and x["C1"]["beg"] == 100, x["C1"]
    assert "S1" not in x                                         # етказиб берувчи мижозлар рўйхатида эмас
base = "?d1=2026-10-01&d2=2026-10-15"
for mode in ("short", "detail", "only"):
    h = c.get("/erp/rep/stock" + base + "&mode=" + mode).get_data(as_text=True)
    assert "ЛДСП Белый" in h and "Остаток товаров" in h, mode
h = c.get("/erp/rep/stock" + base + "&mode=detail&g=G1").get_data(as_text=True)
assert "№ 101" in h and "№ R1" in h
h = c.get("/erp/rep/stock" + base + "&f=F1").get_data(as_text=True)
assert "ЛДСП Белый" in h
def _xl_text(data):
    import io as _io
    from openpyxl import load_workbook
    wb = load_workbook(_io.BytesIO(data))
    return "\n".join(";".join("" if v is None else str(v) for v in row) for ws in wb for row in ws.iter_rows(values_only=True))
r = c.get("/erp/rep/stock" + base + "&csv=1")
assert r.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" and "ЛДСП Белый" in _xl_text(r.data) and "Lumberture" in _xl_text(r.data)
for mode in ("docs", "lines", "goods", "clients"):
    h = c.get("/erp/rep/sales" + base + "&mode=" + mode).get_data(as_text=True)
    assert "Отчет по накладным" in h and ("Мижоз Бир" in h or "ЛДСП" in h), mode
h = c.get("/erp/rep/sales" + base + "&pay=unpaid").get_data(as_text=True)
assert ">101<" in h and ">102<" not in h, "не расч"
h = c.get("/erp/rep/sales" + base + "&pay=paid").get_data(as_text=True)
assert ">102<" in h and ">101<" not in h, "расч"
h = c.get("/erp/rep/sales" + base + "&c=икки").get_data(as_text=True)
assert ">102<" in h and ">101<" not in h
assert c.get("/erp/rep/sales" + base + "&mode=goods&csv=1").mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
for _u in ("/erp/rep/mutual" + base + "&csv=1", "/erp/rep/mutual" + base + "&cx=C1&mode=detail&csv=1", "/erp/rep/sales" + base + "&mode=clients&csv=1", "/erp/rep/receipts" + base + "&mode=lines&csv=1", "/erp/rep/stock" + base + "&mode=detail&g=G1&csv=1"):
    _r = c.get(_u); assert _r.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", _u; _xl_text(_r.data)
h = c.get("/erp/rep/receipts" + base + "&mode=lines&g=G1").get_data(as_text=True)
assert "Эман база" in h and "R1" in h
h = c.get("/erp/rep/mutual" + base).get_data(as_text=True)
assert "Взаиморасчеты" in h and "Мижоз Бир" not in h                 # контрагент танланмаса — рўйхат чиқмайди, фақат жами
for mode in ("short", "detail", "bal"):
    h = c.get("/erp/rep/mutual" + base + "&cx=C1&mode=" + mode).get_data(as_text=True)
    assert "Мижоз Бир" in h and "Взаиморасчеты" in h, mode
h = c.get("/erp/rep/mutual" + base + "&ck=supplier&cx=C1").get_data(as_text=True)
assert "<td><a" not in h or "Мижоз Бир</a>" not in h
h = c.get("/").get_data(as_text=True)
for u in ("/erp/rep/stock", "/erp/rep/sales", "/erp/rep/receipts", "/erp/rep/mutual"):
    assert u in h, u
print("OK reports")
